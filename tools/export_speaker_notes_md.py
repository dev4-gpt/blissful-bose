from pptx import Presentation
from pathlib import Path

prs = Presentation("/Users/aryamandev/Developer/blissful-bose/presentations/01_CURRENT_SEMINAR/Cross_Modal_Signal_Conflict_Seminar_OptionB.pptx")
out_md = Path("/Users/aryamandev/Developer/blissful-bose/presentations/01_CURRENT_SEMINAR/Cross_Modal_Signal_Conflict_Speaker_Notes.md")
out_legacy_md = Path("/Users/aryamandev/Developer/blissful-bose/presentation/VLM_Cross_Modal_Conflict_OptionB_Speaker_Notes.md")

lines = [
    "# Master Seminar Speaker Script & Defense Guide",
    "**Project:** Cross-Modal Signal Conflict in Post-Trained Vision-Language Models  ",
    "**Subtitle:** Causal Dynamics of Visual Injections and Intentional Context  ",
    "**Presenter:** Aryaman Singh Dev | **Advisor:** Prof. Thao Minh Le  ",
    "**Department:** Computer Science & Engineering, Pennsylvania State University  ",
    "**Date:** October 2026 | **Deck:** `presentations/01_CURRENT_SEMINAR/Cross_Modal_Signal_Conflict_Seminar_OptionB.pptx` (20 Slides)  \n",
    "---\n",
    "## Seminar Overview & Strategy",
    "- **Option B Mandate:** Zero training runs, zero defense template engineering, zero large-scale benchmark curation.",
    "- **Core Scientific Focus:** Empirical causal observation of frozen post-trained weights under counterfactual signal injections.",
    "- **Structure:** 20 slides split into Part I (What Has Been Done — 9 Papers) and Part II (What Else Can Be Done — Causal Probing Methodology).\n",
    "---\n",
    "## Slide-by-Slide Complete Speaker Script & Visual Guide\n"
]

for i, s in enumerate(prs.slides):
    title = ""
    section = ""
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            txts = [l.strip() for l in sh.text_frame.text.splitlines() if l.strip()]
            if not section and txts:
                section = txts[0]
            elif not title and txts:
                title = txts[0]
    pics = [sh for sh in s.shapes if sh.shape_type == 13]
    notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else "No notes provided."
    
    lines.append(f"### Slide {i+1:02d}: {title if title else section}")
    lines.append(f"**Section:** `{section}` | **Visual Assets:** {len(pics)} picture(s) attached\n")
    lines.append("#### Spoken Script / Delivery:")
    formatted_notes = "\n".join([f"> {nl}" for nl in notes.splitlines() if nl.strip()])
    lines.append(formatted_notes if formatted_notes else "> No notes provided.")
    lines.append("\n---\n")

md_content = "\n".join(lines)
out_md.write_text(md_content, encoding="utf-8")
out_legacy_md.write_text(md_content, encoding="utf-8")
print("Wrote speaker notes to:")
print(" -", out_md)
print(" -", out_legacy_md)
