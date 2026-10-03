from pathlib import Path

from pptx import Presentation


DECK = Path(__file__).resolve().parents[1] / "VLM_Safety_Alignment_10_Paper_Proposal.pptx"

NOTES = {
    7: """PAPER 1 — FIGSTEP [1]. Fig. 2 is the attack method: paraphrase a harmful request, render it as typography in an image, then pair it with benign incitement text. Table 1 reports mean ASR of 82.50% for FigStep versus 44.80% for direct text across the paper's six open VLMs. The authors tried each item five times and counted success if any try worked. This motivates our fixed OCR/visual-text attack plus matched text-only and benign-OCR controls; it is not a result on today's models.

PAPER 2 — POLYJAILBREAK [4]. Fig. 6 shows the adaptive loop: profile a black-box target, initialize an attack, then use feedback to refine text/image strategies. Table VI reports an average 83.34% ASR and 3.976/5 harmfulness over eight targets, which are separate metrics. The paper includes GPT-4.1 by that exact name and reports up to 15 optimization steps; that is not necessarily 15 total system calls. In our project it motivates a separately budgeted adaptive tier, success-versus-query curves, and separate accounting for target, agent, judge, and image-generation calls. These papers are not a shared leaderboard.""",
    8: """PAPER 3 — VLSBENCH [2]. Figure 1 illustrates visual safety-information leakage: the text can reveal the harmful intent, so a model may appear safe without relying on the image. Table 4 contains the quantitative alignment comparison. For LLaVA-1.5-7B, total safety (refusal plus warning) is 21.26% for multimodal SFT, 27.01% for multimodal DPO, and 13.99% for each textual baseline. These are not ASR values. We use this paper to motivate leak-resistant test cases and matched original-image, image-removed, masked, and decoy controls.

PAPER 4 — USB [3]. Figure 3 summarizes benchmark construction; Table 2 reports model outcomes across four image/text risk states. The paper's exact model string is Claude-Sonnet4; it reports SR 91.16% and RR 18.30% ± 0.75. SR is safe response on harmful inputs; RR is refusal on harmless inputs, so call it benign-input over-refusal. Higher SR and lower benign RR are desirable. USB contributes broad risk/modality coverage and joint diagnosis, not an alignment method. We will choose and report relevant strata rather than collapse all cells into one score.""",
    9: """PAPER 5 — VLGuard [5]. Figure 3 shows how safety and helpfulness move across fine-tuning choices. Table 2 is the main comparison; keep the column headings visible because it mixes attack rates and safe/helpful conditions. The harmfulness columns include AdvBench vanilla/suffix, XSTest Unsafe, VLGuard Safe-Unsafe/Unsafe, and FigStep; lower is better there. XSTest Safe and VLGuard Safe-Safe are benign/helpfulness indicators; higher is better. For LLaVA-v1.5-7B, FigStep ASR falls from 72.62% at baseline to 0.23% after post-hoc tuning and 0.90% after mixed tuning. That single column does not prove utility was preserved. We use VLGuard as a multimodal SFT baseline, retain matched helpfulness data, and evaluate benign near-neighbors.

PAPER 6 — SPA-VL [6]. Figure 1 shows the data path from image collection to easy/hard questions and preference construction. Each record pairs a question and image with chosen and rejected responses. Table 2 reports DPO/PPO results: the SPA-VL-DPO row gives 0.60% average MM-SafetyBench ASR and 0.00/0.00 on AdvBench vanilla/suffix under that paper's setup. DPO learns from chosen-versus-rejected pairs relative to a reference policy; PPO optimizes a policy using a reward signal and is not another name for the dataset. Table 2's benchmark columns have distinct definitions, so don't merge them. The abstract says 13 categories but section 3.1 says 15 secondary categories; we retain that source discrepancy. This paper informs preference provenance and a standard DPO comparator, not a claim of adaptive-attack generalization.""",
    10: """PAPER 7 — ADPO [7]. Figure 2 explains adversary-aware preference training; Table 1 reports both safety and utility. AR-DPO uses the adversarially trained reference model only; AT-DPO uses the adversary-aware DPO loss only; ADPO combines both. For LLaVA-1.5-7B, base-to-ADPO ASR changes are 64.5 to 5.0 on VisualAdv, 84.0 to 0.5 on MMPGDBlank, 22.2 to 0.0 on MultiTrust Typographic, 55.1 to 0.0 on Multimodal, and 42.0 to 0.2 on Crossmodal. Values are percentages. Utility is separately measured and can move. This is direct prior art, so our contribution must be a controlled transfer/generalization study, not the first adversarial VLM alignment method.

PAPER 8 — HOLISAFE / SAFE-VLM [8]. Figure 2 is the Visual Guard Module architecture; Table 3 is the main HoliSafe-Bench result. The five cases are unsafe image+unsafe text, unsafe image+safe text, safe image+unsafe text, safe image+safe text leading to unsafe output, and safe image+safe text leading to safe output. Only the last is the benign-safe control; the fourth is a compositional-risk test. The cited arXiv version reports Safe-LLaVA-7B at 8.8% mASR and 1.3% RR with Claude-3.5, and 15.3% mASR with GPT-4o. Keep judge and version attached: the result is judge-dependent. Table 1, not Figure 1, is the clearest support for coverage of all five image/text states. We use that state taxonomy as evaluation coverage and treat the guard as one optional comparator, not a component to stack automatically with every training arm.""",
    11: """PAPER 9 — DAVSP [9]. Figure 2 shows the method: a learned visual safety prompt and activation alignment with the base model frozen. Table 2 reports resistance success rate (RSR), not attack success rate. Its cited results are 84.20%/99.20% on FigStep for LLaVA-1.5-13B/Qwen2-VL-7B, and 98.72%/99.12% on MM-SafetyBench SD+TYPO. Higher RSR means greater resistance in that paper. We map this to a distinct visual-side intervention and would measure capability and inference cost; do not compare RSR numerically with another paper's ASR.

PAPER 10 — PRAGMA-VL [10]. Figure 3 lays out staged training and context-aware safety/helpfulness arbitration. Table 2 reports benchmark-specific outcomes. For Qwen2.5-VL-7B, MM-SafetyBench ASR is 31.66% versus 48.75% for its base; SIUO effectiveness/safety is 95.21%/63.47% versus 92.17%/38.78%. Table 3 is needed for any general VQA/math capability claim. The metrics have different meanings, so keep their labels with the values. We treat Pragma-VL as a possible comparator only after checking released artifacts, base-model compatibility, and compute. Neither paper's result is directly comparable to the other.""",
    19: """This is the proposed study design synthesized from the ten-paper review; no cited paper has already validated this exact end-to-end pipeline. The first block separates training sources from evaluation benchmarks. Before splitting, retain provenance, license, source intent, image identity, and a safe counterpart where possible. Split by source intent before making paraphrases, rendered images, generated images, or attack variants. Add exact/perceptual image hashes plus embedding-based near-duplicate review (for example CLIP), and lexical plus semantic-intent checks for text. Similarity thresholds are screening tools, not proof; tune them on development data and adjudicate borderline overlaps without inspecting locked-test outcomes.

The training arms are deliberately separate. The primary comparison is native checkpoint, text-only safety SFT, multimodal safety SFT, and standard multimodal DPO; include ADPO only if its assumptions and recipe are reproducible. For a valid text-only versus multimodal SFT contrast, use matched source intents, risk/benign distribution, examples, split, and training budget. The text-only arm should receive no image by default; a blank image can create an unintended visual training condition. Keep helpfulness data matched. Do not combine all prior interventions into one model.

Weight-updating methods and inference add-ons form useful separate axes, but testing every combination may be too expensive. A small, pre-registered wrapper ablation on the native and best aligned checkpoint is a follow-on option. DAVSP is a learned visual-prompt/activation intervention; Safe-VLM includes an auxiliary visual safety head. Record latency and any extra inference calls.

The locked evaluation combines fixed attacks, adaptive attacks, benchmark strata, benign controls, and image attribution tests. PolyJailbreak's reported 15 optimization steps are not a universal total-call budget. Report target queries and success-versus-query curves separately from auxiliary agent, judge, and image-generation calls/tokens/time/cost. The score block keeps harmful compliance, benign false refusal, answer quality, image grounding, capability, and cost distinct. Exact benchmark versions/counts, scoring rubric, judge audit, confidence intervals, and feasibility gates live in the protocol/runbook and companion evidence guide. The diagram explains the stages and data flow, not every implementation detail or per-paper result. Slides 2 and 3 remain untouched.""",
}


def slide_fingerprint(slide):
    return tuple(
        (shape.shape_type, shape.text if shape.has_text_frame else "")
        for shape in slide.shapes
    )


def main():
    presentation = Presentation(str(DECK))
    if len(presentation.slides) != 24:
        raise RuntimeError(f"Expected 24 slides; found {len(presentation.slides)}")

    protected = {index: slide_fingerprint(presentation.slides[index - 1]) for index in (2, 3)}
    for index, note_text in NOTES.items():
        slide = presentation.slides[index - 1]
        if index in range(7, 12):
            if "PAPER-BY-PAPER EVIDENCE" not in " ".join(
                shape.text for shape in slide.shapes if shape.has_text_frame
            ):
                raise RuntimeError(f"Unexpected content on slide {index}")
        slide.notes_slide.notes_text_frame.text = note_text

    if any(slide_fingerprint(presentation.slides[index - 1]) != fingerprint for index, fingerprint in protected.items()):
        raise RuntimeError("A protected slide changed during note update")
    presentation.save(str(DECK))

    check = Presentation(str(DECK))
    if len(check.slides) != 24:
        raise RuntimeError("Slide count changed during save")
    if any(slide_fingerprint(check.slides[index - 1]) != fingerprint for index, fingerprint in protected.items()):
        raise RuntimeError("Slide 2 or 3 content changed")
    for index in NOTES:
        if len(check.slides[index - 1].notes_slide.notes_text_frame.text.strip()) < 200:
            raise RuntimeError(f"Speaker notes missing on slide {index}")
    print(f"Updated notes on slides {', '.join(map(str, NOTES))}; slides 2 and 3 preserved.")


if __name__ == "__main__":
    main()
