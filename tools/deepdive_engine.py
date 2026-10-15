"""Shared helpers for the paper deep-dive decks (reuses the plain style of build_proposal_deck_v2).

Adds: PDF figure cropping, outline highlights drawn over a cropped figure at the exact position of
given text in the PDF, native bar charts, and a speaker-notes markdown writer.
"""
import sys
from pathlib import Path

import pymupdf
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
from build_proposal_deck_v2 import (ACCENT, FONT, INK, MUTED, RULE, add_fig, add_table, add_text, arrow, box,  # noqa: E402,F401
                                    est_lines, lab, rgb, style_run)

ROOT = Path(__file__).resolve().parents[1]
PDFS = ROOT / "presentation" / "paper_pdfs"
FIGDIR = ROOT / "presentation" / "dd_figs"
GREYS = ["BDBDBD", "8C8C8C", "5E5E5E", "2B2B2B"]
PML = "{http://schemas.openxmlformats.org/presentationml/2006/main}style"


def extract_crops(crops):
    """crops: key -> (pdf, page, x0, y0, x1, y1) in PDF points. Writes FIGDIR/<key>.png at 3.5x."""
    FIGDIR.mkdir(parents=True, exist_ok=True)
    docs = {}
    for key, (pdf, pg, x0, y0, x1, y1) in crops.items():
        doc = docs.setdefault(pdf, pymupdf.open(PDFS / f"{pdf}.pdf"))
        pix = doc[pg - 1].get_pixmap(matrix=pymupdf.Matrix(3.5, 3.5), clip=pymupdf.Rect(x0, y0, x1, y1))
        Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(FIGDIR / f"{key}.png")


def figure(slide, crops, key, x, y, w, h, caption=None, hl=None):
    """Place crop `key` fitted into (w,h) at native aspect; optionally outline PDF text strings.
    hl: list of (text, occurrence_indices or None for all)."""
    xr, yb = add_fig(slide, FIGDIR / f"{key}.png", x, y, w, h, caption)
    if hl:
        pdf, pg, cx0, cy0, cx1, cy1 = crops[key]
        page = pymupdf.open(PDFS / f"{pdf}.pdf")[pg - 1]
        pw, ph = xr - x, yb - y
        for term, idx in hl:
            if isinstance(term, tuple):
                hits = [pymupdf.Rect(*term)]
            else:  # whole-word matches first, so '8.20%' does not match inside '88.20%'
                clip = pymupdf.Rect(cx0, cy0, cx1, cy1)
                hits = [pymupdf.Rect(w[:4]) for w in page.get_text("words", clip=clip) if w[4] == term]
                hits = hits or page.search_for(term, clip=clip)
            for i, r in enumerate(hits):
                if idx is not None and i not in idx:
                    continue
                sx = x + (r.x0 - cx0) / (cx1 - cx0) * pw - 0.03
                sy = y + (r.y0 - cy0) / (cy1 - cy0) * ph - 0.02
                sw = (r.x1 - r.x0) / (cx1 - cx0) * pw + 0.06
                sh_ = (r.y1 - r.y0) / (cy1 - cy0) * ph + 0.04
                shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(sx), Inches(sy), Inches(sw), Inches(sh_))
                shp.fill.background()
                shp.line.color.rgb = rgb(ACCENT)
                shp.line.width = Pt(2)
                shp.shadow.inherit = False
                st = shp._element.find(PML)
                if st is not None:
                    shp._element.remove(st)
    return xr, yb


def bar_chart(slide, x, y, w, h, cats, series, horizontal=True, fmt="0.0", highlight=None, size=11, vmax=None):
    """series: list of (name, values). Single series -> accent bars (non-highlighted grey if highlight given);
    several series -> greys."""
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.name, ch.font.size = FONT, Pt(size)
    ch.has_title = False
    ch.has_legend = len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)
    va, ca = ch.value_axis, ch.category_axis
    va.has_major_gridlines = False
    va.visible = False
    va.minimum_scale = 0
    if vmax:
        va.maximum_scale = vmax
    ca.format.line.color.rgb = rgb(RULE)
    ca.tick_labels.font.size = Pt(size)
    if horizontal:
        ca.reverse_order = True
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(size)
    for si, s in enumerate(plot.series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(ACCENT if len(series) == 1 else GREYS[si % 4])
        if len(series) == 1 and highlight is not None:
            for i in range(len(cats)):
                pt = s.points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(ACCENT if i in highlight else "A9A9A9")
    return gf


def write_notes(path, heading, intro, titles, notes):
    lines = [f"# {heading}", "", intro, ""]
    for num, title, key in titles:
        lines += [f"## Slide {num}: {title}", "", notes[key], ""]
    Path(path).write_text("\n".join(lines))
