from pptx import Presentation
from pathlib import Path

prs = Presentation("/Users/aryamandev/Developer/blissful-bose/presentations/01_CURRENT_RESEARCH_PROPOSAL/Cross_Modal_Signal_Conflict_Research_Proposal.pptx")
out_md = Path("/Users/aryamandev/Developer/blissful-bose/presentations/01_CURRENT_RESEARCH_PROPOSAL/Cross_Modal_Signal_Conflict_Research_Proposal_Notes.md")

lines = [
    "# Academic Research Proposal Defense Guide & Speaker Script",
    "**Project Title:** Cross-Modal Signal Conflict in Post-Trained Vision-Language Models  ",
    "**Subtitle:** Causal Dynamics of Visual Injections and Intentional Context  ",
    "**Researcher / Presenter:** Aryaman Singh Dev | **Advisor:** Prof. Thao Minh Le  ",
    "**Department:** Department of Computer Science & Engineering, Pennsylvania State University  ",
    "**Presentation Type:** Academic Research Proposal (Targeting Conference Submission)  ",
    "**Deck Location:** `presentations/01_CURRENT_RESEARCH_PROPOSAL/Cross_Modal_Signal_Conflict_Research_Proposal.pptx` (20 Slides)  \n",
    "---\n",
    "## Research Proposal Executive Summary",
    "- **Core Problem:** Post-trained VLMs splice unaligned vision encoders onto safety-aligned LLMs, creating catastrophic cross-modal signal conflict when visual and textual inputs contradict each other.",
    "- **Methodological Stance (Advisor Directives):**",
    "  1. NO defense template engineering (no wrapper architectures or input heuristics).",
    "  2. NO benchmark creation (no large-scale uncurated datasets).",
    "  3. NO model training (zero backpropagation; 100% frozen inference probing).",
    "- **Core Proposal:** A counterfactual causal intervention protocol crossing Signal Context (visual saliency sweeps: 12–72pt font, contrast ratios) with Intentional Context (instruction hierarchy levels L0–L3) to map continuous arbitration thresholds and layer-wise attention dynamics.\n",
    "---\n",
    "## Slide-by-Slide Defense Script & Visual Guide\n"
]

for i, s in enumerate(prs.slides):
    title = ""
    section = ""
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            txts = [l.strip() for l in sh.text_frame.text.splitlines() if l.strip()]
            if not section and txts:
                section = txts[0]
            elif not title and txts:
                title = txts[0]
    pics = [sh for sh in s.shapes if sh.shape_type == 13]
    notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else "No notes provided."
    
    lines.append(f"### Slide {i+1:02d}: {title if title else section}")
    lines.append(f"**Proposal Section:** `{section}` | **Attached Figures/Tables:** {len(pics)} picture(s)\n")
    lines.append("#### Spoken Defense Transcript:")
    formatted_notes = "\n".join([f"> {nl}" for nl in notes.splitlines() if nl.strip()])
    lines.append(formatted_notes if formatted_notes else "> No notes provided.")
    lines.append("\n---\n")

md_content = "\n".join(lines)
out_md.write_text(md_content, encoding="utf-8")
print("Wrote proposal notes to:", out_md)
