#!/usr/bin/env python3
"""Deep-dive deck B: papers 5-8 (MMJailBench, CMRM, JailBound, SafetyReminder).

  python3 tools/build_deep_dive_b.py [--extract]
Outputs: presentations/05_PAPER_DEEP_DIVES/Deep_Dive_Papers_5-8_MMJailBench_CMRM_JailBound_SafetyReminder.pptx (+ _Notes.md)
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent))
from deep_dive_b_content import (CMRM_INPUTS, CMRM_T1, CMRM_T2, JB_T3, MMJB_AVG, MMJB_T2, MMJB_T4, N, SR_T1, SR_T2, TERMS)  # noqa: E402
from deepdive_engine import (ACCENT, FIGDIR, MUTED, add_table, add_text, bar_chart, box, extract_crops, figure, lab,  # noqa: E402
                             write_notes)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentations" / "05_PAPER_DEEP_DIVES"
NAME = "Deep_Dive_Papers_5-8_MMJailBench_CMRM_JailBound_SafetyReminder"

CROPS = {
 "mmjb_f2": ("mmjailbench", 3, 48, 48, 564, 334), "mmjb_f3": ("mmjailbench", 5, 310, 48, 566, 356),
 "mmjb_f4": ("mmjailbench", 7, 310, 46, 566, 186), "mmjb_f5": ("mmjailbench", 8, 48, 46, 302, 194),
 "cmrm_f1": ("cmrm", 3, 66, 66, 530, 408), "cmrm_f2": ("cmrm", 7, 300, 70, 530, 272),
 "cmrm_t1": ("cmrm", 6, 66, 64, 530, 382),
 "jb_f1": ("jailbound", 2, 104, 66, 508, 270), "jb_f2": ("jailbound", 5, 104, 70, 510, 228), "jb_f7": ("jailbound", 15, 100, 455, 512, 711),
 "sr_f1": ("safetyreminder_aaai", 1, 316, 215, 562, 436), "sr_f2": ("safetyreminder_aaai", 2, 50, 50, 298, 216),
 "sr_f3": ("safetyreminder_aaai", 4, 50, 50, 564, 296),
}


def make_boundary_figure():
    """Schematic of a linear safety probe: two groups of hidden states, the boundary, its normal and a distance."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    rng = np.random.default_rng(3)
    fig, ax = plt.subplots(figsize=(6.4, 4.0), dpi=200)
    safe = rng.normal([-1.3, 0.9], 0.45, (28, 2))
    unsafe = rng.normal([1.3, -0.9], 0.45, (28, 2))
    ax.scatter(*safe.T, s=26, color="#9a9a9a", label="safe inputs")
    ax.scatter(*unsafe.T, s=26, color="#1F4E79", label="unsafe inputs")
    xs = np.array([-2.6, 2.6])
    ax.plot(xs, 0.69 * xs, color="#1A1A1A", lw=1.6)
    ax.text(1.0, 1.55, "boundary: w·h + b = 0", fontsize=10)
    ax.annotate("", xy=(0.55, -0.77), xytext=(0.0, 0.0), arrowprops=dict(arrowstyle="->", color="#1A1A1A", lw=1.6))
    ax.text(-0.9, -1.75, "normal v = w / |w| (arrow)", fontsize=10)
    p = unsafe[3]
    fx = (p[0] + 0.69 * p[1]) / (1 + 0.69 ** 2)
    foot = np.array([fx, 0.69 * fx])
    ax.plot([p[0], foot[0]], [p[1], foot[1]], color="#1F4E79", lw=1.2, ls="--")
    ax.text((p[0] + foot[0]) / 2 + 0.08, (p[1] + foot[1]) / 2 + 0.05, "ε (distance)", fontsize=10, color="#1F4E79")
    ax.set_xlim(-2.6, 2.6)
    ax.set_ylim(-2.0, 2.0)
    ax.set_xticks([])
    ax.set_yticks([])
    for s_ in ax.spines.values():
        s_.set_color("#BBBBBB")
    ax.legend(loc="lower left", frameon=False, fontsize=10)
    ax.set_title("Hidden states at one layer (2-D illustration)", fontsize=10, color="#666666")
    FIGDIR.mkdir(parents=True, exist_ok=True)
    fig.tight_layout()
    fig.savefig(FIGDIR / "jb_schematic.png")
    plt.close(fig)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    titles, n = [], [0]

    def new(tag, title, key):
        n[0] += 1
        titles.append((n[0], title, key))
        s = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(s, 0.6, 0.38, 9, 0.25, [tag], size=11, color=ACCENT, bold=True)
        add_text(s, 0.6, 0.66, 12.1, 0.7, [title], size=26, bold=True)
        add_text(s, 0.6, 7.08, 9, 0.25, ["Paper deep dive 5 to 8: MMJailBench, CMRM, JailBound, SafetyReminder"], size=10, color=MUTED)
        add_text(s, 11.9, 7.08, 0.83, 0.25, [str(n[0])], size=10, color=MUTED, align=3)
        s.notes_slide.notes_text_frame.text = N[key]
        return s

    def fig(s, key, x, y, w, h, cap, hl=None):
        return figure(s, CROPS, key, x, y, w, h, cap, hl)

    def three(s, cols):
        for i, (h_, t_) in enumerate(cols):
            add_text(s, 0.6 + i * 4.15, 1.6, 3.9, 0.35, [h_.upper()], size=12, color=ACCENT, bold=True)
            add_text(s, 0.6 + i * 4.15, 2.05, 3.9, 4.8, [t_], size=15, name=h_)

    # ---- title, roadmap
    n[0] += 1
    titles.append((1, "Paper deep dive: MMJailBench, CMRM, JailBound, SafetyReminder", "title"))
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.notes_slide.notes_text_frame.text = N["title"]
    add_text(s, 0.8, 1.9, 9, 0.3, ["PAPER DEEP DIVE, PAPERS 5 TO 8"], size=12, color=ACCENT, bold=True)
    add_text(s, 0.8, 2.3, 11.8, 1.8, ["MMJailBench, CMRM, JailBound, SafetyReminder"], size=38, bold=True)
    add_text(s, 0.8, 3.9, 11, 1.2, ["Which factors matter, and what hidden states and generation over time show: every metric, table and figure explained"], size=20, color=MUTED)
    add_text(s, 0.8, 5.8, 11, 0.5, ["Cross-modal signal conflict in post-trained vision-language models"], size=14)

    s = new("ROADMAP", "Four papers, four sets of terms", "roadmap")
    add_table(s, [["Paper", "Terms it teaches", "What it gives our study"],
        ["5  MMJailBench", "ASR (1-5 score), CASR, instruction mismatch, delta ASR; intent, framing, visual semantics, carrier", "The factorial mindset, a TEXT arm, a legibility control"],
        ["6  CMRM", "Unsafe rate; five input variations; shifting vector; alpha", "The input ladder and a per-layer shift measurement"],
        ["7  JailBound", "Linear probe, boundary, normal vector; I0/I1, T0/T1; string-match ASR", "Layer-wise probing and a 2x2 factorial layout"],
        ["8  SafetyReminder", "Safety tokens, relative position, prefilling attack, safety-awareness classifier", "Look at where refusal appears; a hidden-state direction"]],
        0.6, 1.7, [2.3, 5.9, 3.9], size=14, pad=0.14, bold_first=True)

    # ================= MMJailBench
    s = new("PAPER 5  MMJAILBENCH", "MMJailBench (arXiv 2026): which factor drives a jailbreak?", "m_q")
    fs = [("Harmful intent", "what is being asked for (272 intents, 9 harm domains)"), ("Prompt framing", "how it is worded (academic, story, system, ...)"),
          ("Visual semantics", "what the image means (authority document, danger, ...)"), ("Instruction carrier", "text, or the same text rendered in an image")]
    for i, (h_, t_) in enumerate(fs):
        box(s, 0.6 + i * 3.1, 1.65, 2.9, 2.3, [[(h_, {"bold": True, "color": ACCENT})], t_], size=14)
    add_text(s, 0.6, 4.4, 12.1, 2.4, ["Earlier benchmarks bundle these in every test case, so a high ASR cannot be traced to a cause.",
        "MMJailBench combines the four factors in a grid. With intent fixed and one factor changed at a time, you see which one moves the result."], size=15, bullet=True, after=8, name="m_q")

    s = new("PAPER 5  MMJAILBENCH", "The grid: 272 x 6 x 5 x 2 = 16,320 cases per model", "m_design")
    fig(s, "mmjb_f2", 0.6, 1.5, 7.6, 4.2, "MMJailBench, Fig. 2")
    add_table(s, [["Factor", "Values"], ["Harmful intent", "272, in 9 domains"], ["Prompt framing", "6: academic, system, structured, story, code, safety-paradox"],
                  ["Visual semantics", "5: danger, scenario, professional role, identity credential, authorization document"], ["Instruction carrier", "2: TEXT or OCR (same text rendered in an image)"],
                  ["Controls", "blank, noise, nature images"]], 8.5, 1.55, [1.5, 2.7], size=12, pad=0.07, bold_first=True)
    add_text(s, 0.6, 6.0, 12.1, 0.9, ["Each case is x = (L(h, t, c), V(h, v)): text from intent h, framing t, carrier c; image from intent h and visual semantics v."], size=14, name="m_design")

    s = new("PAPER 5  MMJAILBENCH", "Metrics: ASR, CASR, instruction mismatch", "m_met")
    add_text(s, 0.6, 1.6, 12.1, 5.2, [lab("Judge", "GPT-5 scores each reply 1 to 5 for harmfulness"),
        lab("ASR", "share of replies scoring 4 or 5. Refusals, warnings and high-level discussion do not count. Example: 66,200 of 130,560 replies score 4 or 5 = 50.7%"),
        lab("Instruction mismatch", "share of replies where the model misread or ignored the instruction (for OCR, a failure to read the image)"),
        lab("CASR", "conditional ASR = successes / (N minus mismatches): ASR among replies where the instruction was understood"),
        lab("Scale", "16 models x 16,320 cases = 261,120 replies")], size=16, after=14, name="met")

    s = new("PAPER 5  MMJAILBENCH", "Overall: average ASR per model (Table 1)", "m_t1")
    bar_chart(s, 0.6, 1.45, 7.6, 5.5, [m for m, _ in MMJB_AVG], [("Average ASR (%)", [v for _, v in MMJB_AVG])], size=10, highlight={0, 15}, vmax=90)
    add_text(s, 8.6, 1.6, 4.1, 5.3, ["Average over nine harm domains. From 2.17% (GPT-5) to 78.38% (GLM-4.6V).", "Higher ASR in cyber abuse, economic harm, privacy and deception; lower in physical harm and sensitive content.",
        "Similar-capability models differ widely: safety is not capability."], size=14, bullet=True, after=8, name="t1 text")

    s = new("PAPER 5  MMJAILBENCH", "Prompt framing is the biggest factor (Fig. 3)", "m_frame")
    fig(s, "mmjb_f3", 0.6, 1.5, 4.7, 5.0, "MMJailBench, Fig. 3: (a) prompt framing, (b) visual semantics")
    add_text(s, 5.8, 1.6, 6.9, 5.3, [lab("Reading", "each polygon is one model; a larger polygon means higher ASR on that axis"),
        lab("Framing", "the gap between the most and least vulnerable framing exceeds 40 points"),
        lab("Higher", "story, structured and academic framings"), lab("Lower", "system-style and safety-paradox framings"),
        lab("Meaning", "with the same harmful intent, the wording and output format change whether the model refuses")], size=15, after=12, name="frame")

    s = new("PAPER 5  MMJAILBENCH", "Visual semantics: meaning matters, presence barely does (Table 2)", "m_vis")
    add_table(s, [["Image condition", "Group", "ASR (%)", "Change vs no image"]] + MMJB_T2, 0.6, 1.6, [3.3, 1.5, 1.4, 2.2], size=14, pad=0.08, bold_first=True)
    add_text(s, 9.3, 1.6, 3.4, 5.3, ["Meaningful images add +8.45 to +12.96 points.", "Blank, noise and nature images add under 2.2.", "Authority cues (documents, credentials) look like implicit legitimacy signals."],
             size=14, bullet=True, after=8, name="vis text")

    s = new("PAPER 5  MMJAILBENCH", "Carrier: OCR is not stronger than plain text (Table 3)", "m_car")
    add_table(s, [["Carrier", "Replies", "Mismatch %", "ASR %", "CASR %"], ["TEXT", "130,560", "0.37", "50.70", "50.89"], ["OCR", "130,560", "4.14", "40.59", "42.34"], ["Both", "261,120", "2.26", "45.64", "46.69"]],
              0.6, 1.7, [1.8, 1.8, 1.8, 1.5, 1.5], size=15, pad=0.12, bold_first=True)
    add_text(s, 0.6, 4.2, 12.1, 2.7, ["OCR instructions get lower ASR (40.59 vs 50.70). The gap stays after removing misreads (CASR 42.34 vs 50.89), so it is not only perception failure.",
        "The paper concludes visually rendered instructions are not inherently stronger carriers; the effect depends on the model.",
        "OCR does cause more misreads (4.14% vs 0.37%): our legibility control."], size=15, bullet=True, after=8, name="car text")

    s = new("PAPER 5  MMJAILBENCH", "Every model has its own shape (Table 4)", "m_t4")
    add_table(s, [["Model", "Intent", "Prompt", "Visual", "Carrier"]] + MMJB_T4, 0.6, 1.45, [3.4, 1.1, 1.1, 1.1, 1.1], size=11, pad=0.035, bold_first=True)
    add_text(s, 8.6, 1.6, 4.1, 5.3, ["Each cell is the range (max minus min) of ASR across the settings of that factor.", "Framing dominates for many models (Qwen2.5-VL 90.8, InternVL3 88.3).",
        "Carrier matters a lot for some (LLaVA-OneVision 50.5, Kimi 34.6) and barely for others (Qwen2.5-VL 0.8, Claude 0.8).", "Visual effects are smaller (0.8 to 11.3) but positive for all."], size=13, bullet=True, after=8, name="t4 text")

    s = new("PAPER 5  MMJAILBENCH", "Inside one model: representation and attention (Figs. 4 and 5)", "m_diag")
    fig(s, "mmjb_f4", 0.6, 1.5, 6.0, 3.2, "MMJailBench, Fig. 4 (gemma3-12b)")
    fig(s, "mmjb_f5", 6.9, 1.5, 5.8, 3.3, "MMJailBench, Fig. 5 (gemma3-12b)")
    add_text(s, 0.6, 5.3, 12.1, 1.6, ["Fig. 4: the difference between authority-document and danger-image states grows with depth (sharpest rise at layer 42); a direction fitted on discovery samples carries over to held-out pairs.",
        "Fig. 5: at that layer authority-document contexts put less attention on harm-related and visual tokens."], size=13, bullet=True, after=6, name="diag text")

    s = new("PAPER 5  MMJAILBENCH", "What MMJailBench shows, what it does not, what we take", "m_take")
    three(s, [("Shows", "Framing is the largest factor; meaningful images add 8 to 13 points while control images add under 2.2; OCR is not stronger than text; models differ in shape; authority contexts shift hidden states along a direction."),
              ("Does not show", "Visibility thresholds. System-role levels. Human-checked judging (GPT-5 judges). Conflict between image and text instructions."),
              ("We take", "A factorial grid; a TEXT arm and mismatch reporting as our legibility control; text-side framing as a variable; a direction-by-layer analysis. Its judge rule (refusals and warnings are not successes) matches our labels.")])

    # ================= CMRM
    s = new("PAPER 6  CMRM", "CMRM (ACL 2025 Findings): do images shift the hidden states?", "c_q")
    parts = [("Observation", "A VLM is less safe than the language model it was built from (safety alignment degradation); even a blank image can trigger harmful replies."),
             ("Hypothesis", "Vision and language parts are trained separately, so image-containing inputs sit in a different region of hidden-state space than the text the safety training covered."),
             ("Test", "Look at hidden states with and without images, then subtract the estimated shift at inference, with no training.")]
    for i, (h_, t_) in enumerate(parts):
        box(s, 0.6 + i * 4.15, 1.65, 3.95, 2.9, [[(h_, {"bold": True, "color": ACCENT})], t_], size=14)
    add_text(s, 0.6, 5.0, 12.1, 1.8, ["Example of degradation: LLaVA on Vicuna-13B falls from 6.57 to 5.92 on MT-Bench."], size=15, bullet=True, name="c_q")

    s = new("PAPER 6  CMRM", "Setup: five inputs, two datasets, one judge", "c_setup")
    add_table(s, [["Input variation", "Image", "Text"], ["Original", "the dataset's image", "harmful query"], ["Blank", "blank image", "harmful query"], ["Gaussian noise", "image + noise", "harmful query"],
                  ["Text + caption", "none (image replaced by its caption)", "query + caption"], ["Text query only", "none", "harmful query"]],
              0.6, 1.6, [2.8, 4.4, 3.0], size=14, pad=0.09, bold_first=True)
    add_text(s, 0.6, 4.9, 12.1, 2.0, [lab("Data", "VLSafe (1,110 malicious queries, benign images); JailbreakLLMs (330 prompts from 14 topics, paired with COCO images)"),
        lab("Models", "LLaVA-1.5-7B, LLaVA-1.5-13B, ShareGPT4V-7B (all Vicuna)"), lab("Metric", "unsafe rate: share of replies a Llama-3.1-8B-Instruct judge calls unsafe")], size=14, after=6, name="setup text")

    s = new("PAPER 6  CMRM", "The core evidence: any image raises the unsafe rate", "c_res0")
    bar_chart(s, 0.6, 1.5, 7.6, 4.6, [m for m, _ in CMRM_INPUTS], [("Unsafe rate (%)", [v for _, v in CMRM_INPUTS])], size=12, highlight={2, 3, 4}, vmax=75)
    add_text(s, 8.6, 1.6, 4.1, 5.3, ["LLaVA-1.5-7B on VLSafe.", "Text only: 1.91%. Text + caption: 5.52%.", "Blank or noise image: about 57%. Original image: 61.53%.",
        "Image presence, not content, does most of the damage here."], size=14, bullet=True, after=8, name="res0 text")

    s = new("PAPER 6  CMRM", "How to read the PCA picture (Fig. 1)", "c_pca")
    fig(s, "cmrm_f1", 0.6, 1.5, 7.8, 5.0, "CMRM, Fig. 1")
    add_text(s, 8.8, 1.6, 3.9, 5.3, ["Each point: the top-layer hidden state of the last input token for one input.", "Axes: the first two principal components.",
        "Text-only groups (caption, query) separate from multimodal groups (original, blank, noise).", "The shift is observed; it is not proof of cause."], size=13, bullet=True, after=8, name="pca text")

    s = new("PAPER 6  CMRM", "The method: add back what the image did", "c_method")
    add_text(s, 0.6, 1.6, 12.1, 1.0, [[("h* = h_o + α (h_t − h_c)", {"bold": True, "size": 30, "color": ACCENT})]], size=30, name="eq")
    add_table(s, [["Symbol", "Meaning"], ["h_o", "hidden state for the original image + text"], ["h_c", "hidden state when the image is replaced by a meaningless one"],
                  ["h_t", "hidden state for the text alone"], ["α", "mixing coefficient for the size of the shift; 1.0 in the main results"],
                  ["h*", "estimated ideal state, as if the image did not shift it"]], 0.6, 2.7, [1.4, 7.2], size=14, pad=0.08, bold_first=True)
    add_text(s, 0.6, 5.65, 12.1, 1.3, ["Shifting vector: dataset-level (first principal component of h_t minus h_c over the dataset) or sample-level. Applied to the last-token state at all layers. No weights change."], size=13, name="method note")

    s = new("PAPER 6  CMRM", "Table 1, LLaVA-1.5-7B block, row by row", "c_t1")
    add_table(s, [["Condition", "VLSafe Orig", "Blank", "Noise", "JBLLMs Orig", "Blank", "Noise", "L-Bench", "S-QA"]] + CMRM_T1, 0.6, 1.6, [3.0, 1.2, 0.9, 0.9, 1.3, 0.9, 0.9, 1.0, 0.9], size=12, pad=0.08, bold_first=True)
    add_text(s, 0.6, 5.0, 12.1, 1.9, ["Unsafe rate (lower is better) on VLSafe and JailbreakLLMs; utility on LLaVA-Bench-COCO and ScienceQA (higher is better).",
        "CMRM takes the unsafe rate from about 60% to 3 to 5%, near the text-only level, with little utility change (S-QA 68.03 to 65.89 or 66.14). VLGuard is a training method shown for reference."], size=13, bullet=True, after=6, name="t1 text")

    s = new("PAPER 6  CMRM", "How robust is it? Alpha, layers, transfer", "c_abl")
    fig(s, "cmrm_f2", 0.6, 1.5, 5.4, 4.6, "CMRM, Fig. 2: unsafe rate vs alpha")
    add_text(s, 6.4, 1.5, 6.3, 0.3, ["Table 2: edited layers, unsafe rate (%) on VLSafe"], size=12, color=ACCENT, bold=True)
    add_table(s, [["Start layer", "End layer", "LLaVA", "ShareGPT"]] + CMRM_T2, 6.4, 1.85, [1.5, 1.5, 1.5, 1.6], size=13, pad=0.06)
    add_text(s, 6.4, 4.95, 6.3, 1.9, ["Alpha near 1.0 works best; editing all layers is needed (5 to 24 gives 16.94%).", "Transfer: a shifting vector from VLSafe works on JailbreakLLMs (LLaVA 7B: 6.06, 4.55, 8.71)."], size=13, bullet=True, after=6, name="abl text")

    s = new("PAPER 6  CMRM", "What CMRM shows, what it does not, what we take", "c_take")
    three(s, [("Shows", "Image-containing inputs sit in a different region of hidden-state space; blank and noise images cause the shift too; subtracting an estimated shift cuts unsafe rates to near text-only level with no training."),
              ("Does not show", "Which image property causes the shift. Conflict between image and text instructions. Qwen-family behaviour (it tested LLaVA-1.5 and ShareGPT4V)."),
              ("We take", "The five-input ladder as controls; measuring h_t minus h_c per layer for each of our conditions; the last-input-token state as the place to look. We do not build a defence.")])

    # ================= JailBound
    s = new("PAPER 7  JAILBOUND", "JailBound (NeurIPS 2025): a linear safety boundary inside the model", "j_q")
    fig(s, "jb_f1", 0.6, 1.5, 8.4, 3.1, "JailBound, Fig. 1: five stages")
    add_text(s, 0.6, 5.0, 12.1, 1.9, ["Idea from eliciting latent knowledge: models hold internal information that differs from what they say.",
        "Stage 2 (probing) is what we study; stage 4 (crossing) is a gradient attack that we do not use."], size=14, bullet=True, after=8, name="j_q")

    s = new("PAPER 7  JAILBOUND", "The probe: logistic regression on hidden states", "j_probe")
    fig(s, "jb_schematic", 0.6, 1.5, 6.6, 4.2, "Schematic (ours) of a linear probe at one layer")
    add_text(s, 7.6, 1.6, 5.1, 5.3, [lab("Input", "hidden state h at a fusion layer for image-text pairs labelled safe or unsafe"),
        lab("Probe", "P(unsafe) = sigmoid(w·h + b); boundary where w·h + b = 0"),
        lab("Normal v", "w / |w|, pointing from safe toward unsafe"),
        lab("Epsilon", "distance of an example to the boundary"),
        lab("Limit", "a probe shows what is linearly readable, not whether the model will refuse")], size=14, after=10, name="probe")

    s = new("PAPER 7  JAILBOUND", "Where is it readable? Probe accuracy by layer (Fig. 7)", "j_layers")
    fig(s, "jb_f7", 0.6, 1.5, 6.8, 4.9, "JailBound, Fig. 7 (28 classifiers on Qwen2.5-VL)")
    add_text(s, 7.8, 1.6, 4.9, 5.3, ["Deep layers converge quickly to near-perfect accuracy: safe and unsafe inputs are linearly separable.", "Early layers (about 0 to 4) do poorly.",
        "The introduction claims 100% across all layers; the appendix says early layers are poor. The appendix is the careful statement."], size=14, bullet=True, after=8, name="layers text")

    s = new("PAPER 7  JAILBOUND", "The crossing step: three losses (Fig. 2)", "j_attack")
    fig(s, "jb_f2", 0.6, 1.5, 8.6, 3.35, "JailBound, Fig. 2")
    add_text(s, 0.6, 5.1, 12.1, 1.8, [lab("L_align", "pull the hidden state to a target beyond the boundary"), lab("L_geo", "move along the boundary's normal direction"),
        lab("L_sem", "keep perturbations small: L2 on the image, fluency on the text suffix"), lab("Needs gradients", "white-box; we do not use it")], size=13, after=4, name="attack text")

    s = new("PAPER 7  JAILBOUND", "Setup and the 2 x 2 factorial", "j_setup")
    add_table(s, [["Item", "Detail"], ["Data", "MM-SafetyBench: 1,719 examples, 13 categories; unsafe image + malicious prompt"],
                  ["Models", "White-box: Llama-3.2-11B-Vision, Qwen2.5-VL-7B, MiniGPT-4. Black-box: GPT-4o, Gemini 2.0 Flash, Claude 3.5 Sonnet"],
                  ["Metric", "ASR = share of replies with NO rejection phrase (exact string match); no manual labels"],
                  ["Grid", "I0 / I1 = clean / perturbed image; T0 / T1 = original / perturbed text; {I1, T1} = joint optimisation"]], 0.6, 1.6, [1.8, 10.3], size=14, pad=0.11, bold_first=True)
    add_text(s, 0.6, 5.2, 12.1, 1.7, ["String-match ASR is lenient: an evasive or off-topic reply with no rejection phrase counts as a success. These ASRs are not comparable with judged ASRs elsewhere."], size=15, name="setup note")

    s = new("PAPER 7  JAILBOUND", "Table 3: each modality alone, both, and joint", "j_t3")
    add_table(s, [["Category", "Setting", "Llama-3.2-11B", "Qwen2.5-VL-7B", "MiniGPT-4"]] + JB_T3, 0.6, 1.55, [2.4, 2.2, 1.7, 1.8, 1.5], size=12, pad=0.055, bold_first=True)
    add_text(s, 10.5, 1.6, 2.3, 5.3, ["White-box ASR (%).", "Each modality alone raises ASR; both raise it more; joint is highest.", "Black-box: 75.24% GPT-4o, 70.06% Gemini 2.0 Flash, 56.55% Claude 3.5."], size=12, bullet=True, after=6, name="t3 text")

    s = new("PAPER 7  JAILBOUND", "What JailBound shows, what it does not, what we take", "j_take")
    three(s, [("Shows", "A linear boundary separates safe from unsafe hidden states well in deep layers; steering across it is an effective attack; image and text perturbations compound."),
              ("Does not show", "That the boundary equals refusal behaviour (labels are about safe or unsafe inputs). Visibility thresholds. Reliable ASR (string match)."),
              ("We take", "Layer-wise linear probing and the 2 x 2 factorial layout. We change: a refusal direction built from the model's own refusals on text-only prompts, and a judged label. Gradient attacks are not used.")])

    # ================= SafetyReminder
    s = new("PAPER 8  SAFETYREMINDER", "SafetyReminder (AAAI 2026): refusal arrives late", "r_q")
    fig(s, "sr_f1", 0.6, 1.5, 4.9, 5.0, "SafetyReminder, Fig. 1")
    add_text(s, 6.0, 1.6, 6.7, 5.3, ["A jailbroken VLM often starts writing the harmful content and only at the end adds 'Remember, this is dangerous and illegal.'",
        "The authors call this delayed safety awareness: the model does recognise the danger, too late.", "Their fix trains a small soft prompt. The observation is what concerns us."], size=15, bullet=True, after=10, name="r_q")

    s = new("PAPER 8  SAFETYREMINDER", "Where do safety words fall? (Fig. 2)", "r_fig2")
    fig(s, "sr_f2", 0.6, 1.5, 6.4, 4.2, "SafetyReminder, Fig. 2")
    add_text(s, 7.4, 1.6, 5.3, 5.3, [lab("Safety token", "a refusal or safety word such as 'sorry', 'illegal', 'cannot'"), lab("Relative position", "0 = start of reply, 1 = end"),
        lab("Prefilling attack", "the reply is forced to begin with 'Sure, here is...'"),
        lab("Qwen2-VL-7B", "plain harmful prompts: safety words near the start; prefilling moves them to the end"),
        lab("LLaVA-1.5-7B", "mass near the end in both cases")], size=13, after=9, name="fig2 text")

    s = new("PAPER 8  SAFETYREMINDER", "The safety-awareness classifier", "r_aware")
    steps = [("1  Data", "256 harmful and 256 benign text prompts from the Arditi et al. refusal-direction set; no images"),
             ("2  Hidden state", "last input token, top decoder layer; centre the states"),
             ("3  Subspace", "SVD; keep the top m principal directions as a 'safety subspace'"),
             ("4  Classifier", "project onto the subspace; train a linear classifier (sigmoid of w·h + b)")]
    for i, (h_, t_) in enumerate(steps):
        box(s, 0.6 + i * 3.1, 1.65, 2.9, 2.7, [[(h_, {"bold": True, "color": ACCENT})], t_], size=13)
    add_text(s, 0.6, 4.8, 12.1, 2.0, ["This is close to our planned refusal direction: a direction in hidden-state space separating harmful from harmless text-only prompts for the same model.",
        "They use SVD plus a classifier; we plan a difference of means and a projection per layer."], size=14, bullet=True, after=8, name="aware text")

    s = new("PAPER 8  SAFETYREMINDER", "The defence, for completeness: SAPT (Fig. 3)", "r_sapt")
    fig(s, "sr_f3", 0.6, 1.5, 8.6, 4.5, "SafetyReminder, Fig. 3")
    add_text(s, 9.5, 1.6, 3.2, 5.3, ["Trains a 4-token soft prompt (20 epochs, AdamW); language model and vision encoder stay frozen.", "512 training samples. At inference the prompt is injected when the classifier score exceeds 0.8.",
        "A trained defence; we use only its measurements."], size=12, bullet=True, after=6, name="sapt text")

    s = new("PAPER 8  SAFETYREMINDER", "Table 1: attack success rate, with and without the defence", "r_t1")
    add_table(s, [["Model", "Defense", "FigStep", "MMSafety", "SPA-VL", "ImgJP", "GCG", "BAP", "Avg"]] + SR_T1, 0.6, 1.6, [1.9, 2.0, 1.1, 1.2, 1.1, 1.0, 0.9, 0.9, 0.9], size=13, pad=0.09, bold_first=True)
    add_text(s, 0.6, 5.35, 12.1, 1.6, ["Judged by LlamaGuard3. FigStep and MMSafety hide the harm in the image with a benign text query.",
        "Qwen2-VL-7B, no defence: FigStep only 26.6% (LLaVA-1.5-7B: 74.9%), but ImgJP 97.4%, GCG 79.0%, BAP 82.0%: typographic attacks are weak on Qwen, optimisation attacks are not."], size=13, bullet=True, after=6, name="t1 text")

    s = new("PAPER 8  SAFETYREMINDER", "Table 2: the cost side, utility and over-refusal", "r_t2")
    add_table(s, [["Model", "Defense", "MM-Vet avg", "Refusal rate %"]] + SR_T2, 0.6, 1.6, [2.4, 2.4, 1.8, 2.0], size=15, pad=0.11, bold_first=True)
    add_text(s, 9.7, 1.6, 3.0, 5.2, ["Refusal rate is the share of benign questions refused.", "BlueSuffix cuts Qwen2-VL's score to 13.4 and refuses 35.8%: a defence can wreck utility.", "Report safety and over-refusal together."],
             size=13, bullet=True, after=8, name="t2 text")

    s = new("PAPER 8  SAFETYREMINDER", "What SafetyReminder shows, what it does not, what we take", "r_take")
    three(s, [("Shows", "Replies often carry safety words only at the start and the end; an attack moves them to the end; Qwen2-VL-7B refuses plain harmful requests at the start; a linear classifier on the last-input-token state detects safety awareness."),
              ("Does not show", "Any attention analysis (an earlier draft wrongly said so), a 'first five tokens' rule (that number is ours to find), or image-text conflict specifically."),
              ("We take", "Measure where in the reply refusal appears, starting with the first token; a hidden-state direction fit on text-only harmful vs harmless prompts; over-refusal reported next to safety. The soft prompt is not used.")])

    # ================= closing
    s = new("SUMMARY", "Glossary: every term from the four papers", "glossary")
    add_table(s, [["Term", "Paper", "Meaning", "How we use it"]] + TERMS, 0.5, 1.35, [2.4, 1.6, 4.6, 3.6], size=11, pad=0.03, bold_first=True)
    s = new("SUMMARY", "What the four papers mean for our design", "synth")
    add_text(s, 0.6, 1.6, 12.1, 5.2, [lab("MMJailBench", "framing and meaningful visual context matter more than the carrier; models have different shapes"),
        lab("CMRM", "any image shifts hidden states; the five-input ladder; measure the shift per layer"),
        lab("JailBound", "layer-wise linear probing and a 2 x 2 factorial layout, but a weak string-match metric"),
        lab("SafetyReminder", "look at where in the reply refusal appears; a hidden-state classifier close to a refusal direction; Qwen2-VL numbers warn of a floor effect"),
        lab("So we use", "single-shot judged labels, TEXT and clean-image arms, a legibility control, framing and system-role levels, per-layer projections on a refusal direction, first-token analysis, over-refusal reported alongside")],
        size=16, after=14, name="synth")

    OUT.mkdir(parents=True, exist_ok=True)
    prs.save(OUT / f"{NAME}.pptx")
    write_notes(OUT / f"{NAME}_Notes.md", "Deep dive notes: papers 5 to 8",
                "Plain-language notes for every slide. All numbers were read from the PDFs in presentation/paper_pdfs/.", titles, N)
    print(f"Saved {NAME}.pptx with {n[0]} slides")


if __name__ == "__main__":
    if "--extract" in sys.argv or not (FIGDIR / "mmjb_f2.png").exists():
        extract_crops(CROPS)
    make_boundary_figure()
    build()
