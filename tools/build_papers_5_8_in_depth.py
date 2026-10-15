#!/usr/bin/env python3
"""Ten-slide in-depth deck on papers 5-8 (MMJailBench, CMRM, JailBound, SafetyReminder).

  python3 tools/build_papers_5_8_in_depth.py
Outputs: presentations/05_PAPER_DEEP_DIVES/Papers_5-8_In_Depth_10_Slides.pptx (+ _Notes.md)
Figures: crops already in presentation/dd_figs (made by build_deep_dive_b.py). Numbers: deep_dive_b_content.py.
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent))
from deep_dive_b_content import CMRM_INPUTS, JB_T3, MMJB_T2, SR_T1  # noqa: E402
from deepdive_engine import ACCENT, FIGDIR, MUTED, add_fig, add_table, add_text, bar_chart, figure, lab, write_notes  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentations" / "05_PAPER_DEEP_DIVES"
NAME = "Papers_5-8_In_Depth_10_Slides"
FOOT = "Papers 5 to 8 in depth: MMJailBench, CMRM, JailBound, SafetyReminder"
CROPS = {"mmjb_f2": ("mmjailbench", 3, 48, 48, 564, 334), "cmrm_f1": ("cmrm", 3, 66, 66, 530, 408),
         "jb_f7": ("jailbound", 15, 100, 455, 512, 711), "sr_f2": ("safetyreminder_aaai", 2, 50, 50, 298, 216)}

N = {}
N["title"] = """In-depth look at the second group of four papers. Each gets two slides. The goal is that for every number on a slide you can say what was measured, how, and what it means for our design.
Order: MMJailBench (which factor drives jailbreaks), CMRM (do images shift hidden states), JailBound (linear probes inside the model), SafetyReminder (where refusal appears in a reply). The last slide ties the four to our design and lists what we deliberately do not take.
The longer 36-slide deck has every table row by row; this one keeps the argument and the key evidence."""
N["m_design"] = """MMJailBench's question: which part of a multimodal jailbreak does the work? It builds a full grid so each factor can be varied alone. 272 harmful intents across 9 harm domains, 6 prompt framings (academic, system, structured, story, code, safety-paradox), 5 visual-semantic conditions (danger context, scenario, professional role, identity credential, authorization document) and 2 carriers (plain TEXT or OCR image). 272 x 6 x 5 x 2 = 16,320 instances per model; 16 models give 261,120 replies.
Metrics. A GPT-5 judge scores each reply 1 to 5 for harmfulness; ASR is the share scoring 4 or 5. Refusals, warnings and high-level discussion do not count. Example: 66,200 of 130,560 replies scoring 4 or 5 is 50.7%. For OCR there is a risk that the model simply misreads the image, so the paper also reports instruction mismatch (misread or irrelevant replies) and CASR = successes / (N minus mismatches). Delta ASR is the change against the no-image condition.
For us: ASR with a model judge is our kind of outcome; mismatch and CASR are the same idea as our legibility covariate."""
N["m_res"] = """Left, Table 2 (visual semantics): ASR change against no image (34.77% baseline). The five meaningful contexts add +8.45 to +12.96 points, authority documents highest. The controls add little: blank +2.19, noise +0.76, nature +0.72. So it is the meaning of the visual context, not the mere presence of an image, that raises ASR here.
Right, Table 3 (carrier): TEXT 50.70% ASR, OCR 40.59%; CASR TEXT 50.89%, OCR 42.34%; mismatch TEXT 0.37%, OCR 4.14%. OCR is weaker even after removing misreads, so rendering an instruction in an image is not automatically stronger. The long deck's Table 4 shows the carrier effect is model-dependent (about 50 points for LLaVA-OneVision, about 1 for Qwen2.5-VL and Claude).
Third finding, not charted: prompt framing is the largest factor, with more than 40 points between best and worst style.
Lessons: always run a TEXT arm; treat text-side framing as a variable; report misreads. Caveat: ASR only, internals on one model (gemma3-12b)."""
N["c_ev"] = """CMRM's observation: a VLM's language backbone is safe on text, but adding an image degrades safety. Setup: the same harmful text query in five input forms: text only, text plus the image's caption, text plus a Gaussian-noise image, text plus a blank image, and original image plus text. Dataset VLSafe (1,110 malicious queries), model LLaVA-1.5-7B, metric unsafe rate from a Llama-3.1-8B-Instruct judge.
Chart: text only 1.91%, caption 5.52%, noise image 57.21%, blank image 56.87%, original 61.53%. With no image or only a caption the model behaves like its safe backbone; with any image, even a blank one, the rate jumps by more than 50 points. For this model and dataset, the image's presence does most of the damage. That is why our design has a clean-image arm.
Method, for understanding only: h_o is the hidden state for the original input, h_c for a corrupted image, h_t for text alone. Assume the state is the ideal state plus an image-induced shift, which gives h* = h_o + alpha (h_t minus h_c). It is applied at inference with no weight change. We do not build this as a defense; the size of h_t minus h_c per layer is a measurement we reuse."""
N["c_pca"] = """Fig. 1 is a 2-D PCA of last-token hidden states for LLaVA-7B, LLaVA-13B and ShareGPT4V-7B. PCA projects the high-dimensional states onto the two directions of largest variance so clusters become visible. Colours are the five input types. The text-only and caption points form their own cloud; every variant containing an image falls in a separate region, even the blank one, which matches the chart on the previous slide.
Table 1 (LLaVA-1.5-7B), on the right: with no intervention the unsafe rate is 61.53% (original), 56.87% (blank); CMRM sample-level correction brings these to 3.15% and 1.24%. Cost: general capability L-Bench 79.2 to 77.3 and S-QA 68.03 to 66.14, small drops. Training-based VLGuard also reaches near zero but changes weights.
Limits: it shows images shift states, not which image property matters, and the evidence is on LLaVA-family models. For us: compute the shift per condition and layer, relative to a text-only baseline, with the blank-image control."""
N["j_probe"] = """JailBound's first stage is a linear probe. Take image-text examples labelled safe or unsafe, read the hidden state h at a fusion layer, and train a logistic regression: P(unsafe) = sigmoid(w.h + b). The boundary is the plane w.h + b = 0; the unit normal v = w / |w| points from safe toward unsafe; epsilon is an example's distance to the plane. One classifier per layer.
The schematic (ours, not the paper's) shows those pieces. Fig. 7 (appendix A.2) shows per-layer training: deep layers reach near-perfect accuracy while layers 0 to 4 do poorly; the introduction's wording that all layers reach 100% is the looser statement.
What a probe tells us: whether safe-versus-unsafe information is linearly readable at a layer. What it does not: whether the model will refuse, since the labels describe input content ('determined by the VLM' on MM-SafetyBench samples). So we borrow the layer-wise linear method and aim it at refusal behaviour instead."""
N["j_t3"] = """Setup: MM-SafetyBench, 1,719 adversarial examples over 13 prohibited categories; six VLMs (white-box on Llama-3.2-11B, Qwen2.5-VL-7B, MiniGPT-4 among others; black-box on GPT-4o, Gemini 2.0 Flash, Claude 3.5 Sonnet). I0 and I1 are the clean and perturbed image, T0 and T1 the original text and the text with an optimised suffix.
Chart, Illegal Activity, white-box ASR: no attack I0+T0 51.47 / 2.94 / 42.65 (Llama / Qwen / MiniGPT-4); text suffix only 82.35 / 30.88 / 69.12; image only 85.29 / 22.06 / 83.82; both separately 88.24 / 64.71 / 95.59; joint optimisation 95.59 / 82.35 / 100.00. Malware Generation follows the same shape. Read it as a factorial: each channel alone helps, both together help more, joint best.
The metric caveat matters: ASR here is the share of replies with no rejection phrase (exact string match). That counts evasive or gibberish replies as successes, so these numbers are not comparable with judged ASRs elsewhere. Note the baseline: Qwen2.5-VL-7B starts at 2.94%, much better at emitting refusal phrases than Llama-3.2 (51.47%). Abstract averages: 94.32% white-box, 67.28% black-box. The attack needs gradients; we do not use it."""
N["r_fig2"] = """SafetyReminder's observation: a jailbroken VLM often starts to comply and only adds 'this is illegal' at the end. Terms: a safety token is a refusal or warning word ('sorry', 'illegal', 'cannot'); its relative position is where it falls in the reply, 0 at the start and 1 at the end; a prefilling attack forces the reply to begin with an affirmative opening such as 'Sure, here is'.
Fig. 2 plots the density of safety tokens by position, for a harmful dataset and for a prefilling attack. For LLaVA-1.5-7B both curves peak near the end. For Qwen2-VL-7B the plain harmful prompts put safety tokens near the start (the model refuses early) and the prefilling attack moves them to the end. Reading: safety wording clusters at the two ends and is largely absent in the middle, and an attack pushes it to the end.
For us: Qwen2-VL is our model family and it refuses early when not attacked, which supports measuring the first token and the first few tokens rather than only the final answer. Earlier drafts attributed an attention result and a 'first five tokens' rule to this paper; neither is in it."""
N["r_t1"] = """Table 1, ASR (%) on six test sets. LLaVA-1.5-7B with no defense: FigStep 74.9, MMSafety 58.2, SPA-VL 45.6, ImgJP 86.6, GCG 94.0, BAP 97.0, average 76.1. Qwen2-VL-7B no defense: 26.6, 13.2, 1.5, 97.4, 79.0, 82.0, average 50.0. Self-Reminder (a prompt-only defense) helps partly; their trained soft prompt (SAPT) brings the averages to 5.0 and 4.7.
Two readings for us. First, the attack gap between models: FigStep gets 74.9% on LLaVA but 26.6% on Qwen2-VL-7B, which supports the floor-effect risk for our Qwen2-VL-2B pilot, though the 2B model has not been measured. Second, Table 2 reports the cost side: utility on MM-Vet (LLaVA 27.9 to 28.0, Qwen 52.3 to 52.6) and the refusal rate on benign questions (2.3 to 3.7, and 3.7 to 4.1), while another defense, BlueSuffix, cuts Qwen2-VL's capability to 13.4 and lifts refusals to 35.8%. The lesson is to report safety and over-refusal together.
The safety-awareness classifier (SVD plus a linear classifier on the last-token hidden state, over 256 harmful and 256 benign text prompts from the Arditi et al. set) parallels our planned refusal direction; we would use a difference of means per layer. We do not use SAPT, which is trained."""
N["synth"] = """Putting the four together. MMJailBench says meaning and framing matter more than the image being present as text, and that OCR is not automatically stronger. CMRM says any image shifts the hidden state. JailBound says safety information is linearly readable, weakly in early layers. SafetyReminder says refusal has a position in the reply.
What we take: a TEXT arm and a legibility covariate; a clean-image arm and a per-layer shift against a text-only baseline; layer-wise linear analysis aimed at refusal; the first-token and early-token refusal probabilities.
What we do not take: CMRM's correction (a defense), JailBound's gradient attack and string-match ASR, SafetyReminder's trained soft prompt. Three of the four papers contain a method that trains or uses gradients; only their measurement ideas are borrowed.
Still unverified: the refusal-direction work of Arditi et al. is cited from memory, and related typographic sweeps outside these eight papers have not been read in full."""


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    titles, n = [], [0]

    def new(tag, title, key):
        n[0] += 1
        titles.append((n[0], title, key))
        s = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(s, 0.6, 0.38, 9, 0.25, [tag], size=11, color=ACCENT, bold=True)
        add_text(s, 0.6, 0.66, 12.1, 0.7, [title], size=26, bold=True)
        add_text(s, 0.6, 7.08, 9, 0.25, [FOOT], size=10, color=MUTED)
        add_text(s, 11.9, 7.08, 0.83, 0.25, [str(n[0])], size=10, color=MUTED, align=3)
        s.notes_slide.notes_text_frame.text = N[key]
        return s

    # 1 title
    n[0] += 1
    titles.append((1, "Papers 5 to 8 in depth", "title"))
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.notes_slide.notes_text_frame.text = N["title"]
    add_text(s, 0.8, 1.7, 9, 0.3, ["PAPERS 5 TO 8, IN DEPTH"], size=12, color=ACCENT, bold=True)
    add_text(s, 0.8, 2.1, 11.5, 1.2, ["MMJailBench, CMRM, JailBound, SafetyReminder"], size=38, bold=True)
    add_text(s, 0.8, 3.5, 11, 1.0, ["What each paper measured, how, what the numbers mean, and what each gives our study"], size=20, color=MUTED)
    add_table(s, [["Paper", "Terms it teaches"],
        ["MMJailBench", "ASR (GPT-5 judge), instruction mismatch, CASR, delta ASR"], ["CMRM", "unsafe rate, hidden-state shift, PCA, h* = h_o + α(h_t − h_c)"],
        ["JailBound", "linear probe, boundary, normal vector, ε, I0/I1 and T0/T1, string-match ASR"], ["SafetyReminder", "safety token, relative position, prefilling attack, safety-awareness classifier"]],
        0.8, 4.6, [2.4, 8.8], size=13, pad=0.07, bold_first=True)

    # 2 MMJailBench design and metrics
    s = new("MMJAILBENCH (arXiv 2026), 1 OF 2", "A full grid, so each factor can be separated", "m_design")
    xr, yb = figure(s, CROPS, "mmjb_f2", 0.6, 1.5, 7.0, 3.9, "MMJailBench, Fig. 2")
    add_text(s, 0.6, yb + 0.3, 7.0, 1.2, ["272 intents x 6 framings x 5 visual contexts x 2 carriers = 16,320 cases per model; 16 models."], size=14, name="grid")
    add_text(s, 8.0, 1.5, 4.7, 5.4, [lab("ASR", "share of replies a GPT-5 judge scores 4 or 5 out of 5. Refusals and warnings do not count. 66,200 of 130,560 replies is 50.7%."),
        lab("Mismatch", "share of replies where the model misread or ignored the instruction"),
        lab("CASR", "ASR among replies where it was understood: successes / (N minus mismatches)"),
        lab("Delta ASR", "change against the no-image condition")], size=14, after=11, name="metrics")

    # 3 MMJailBench results
    s = new("MMJAILBENCH, 2 OF 2", "Meaning moves ASR; an OCR carrier is weaker than text", "m_res")
    add_text(s, 0.6, 1.45, 6.4, 0.3, ["TABLE 2: ASR CHANGE VS NO IMAGE (POINTS)"], size=12, color=ACCENT, bold=True)
    bar_chart(s, 0.6, 1.8, 6.4, 4.9, [r[0] for r in MMJB_T2], [("Delta ASR", [float(r[3]) for r in MMJB_T2])], fmt="0.00", highlight=[0, 1, 2, 3, 4], size=12, vmax=16)
    yb = add_table(s, [["Carrier", "ASR", "CASR", "Mismatch"], ["TEXT", "50.70", "50.89", "0.37"], ["OCR image", "40.59", "42.34", "4.14"]],
                   7.6, 1.8, [1.8, 1.1, 1.1, 1.1], size=14, pad=0.09, bold_first=True)
    add_text(s, 7.6, yb + 0.25, 5.1, 3.6, ["Dark bars: meaningful visual context (+8.45 to +12.96). Grey: blank, noise, nature and no image (up to +2.19).",
        "OCR stays lower after removing misreads, so rendered instructions are not automatically stronger.",
        "Prompt framing is the biggest factor: over 40 points between best and worst style."], size=13, bullet=True, after=8, name="m res text")

    # 4 CMRM evidence and method
    s = new("CMRM (ACL 2025 FINDINGS), 1 OF 2", "Any image, even a blank one, raises the unsafe rate", "c_ev")
    add_text(s, 0.6, 1.45, 6.4, 0.3, ["LLAVA-1.5-7B ON VLSAFE: UNSAFE RATE (%)"], size=12, color=ACCENT, bold=True)
    bar_chart(s, 0.6, 1.8, 6.4, 4.2, [c for c, _ in CMRM_INPUTS], [("Unsafe rate", [v for _, v in CMRM_INPUTS])], fmt="0.00", highlight=[2, 3, 4], size=12, vmax=75)
    add_text(s, 0.6, 6.1, 6.4, 0.8, ["Text-only inputs stay near the safe backbone; every input with an image jumps by over 50 points."], size=13, color=MUTED, name="c note")
    add_text(s, 7.5, 1.5, 5.2, 5.4, [lab("h_o", "hidden state for the original image + text"), lab("h_c", "state with a meaningless (corrupted) image"), lab("h_t", "state for the text alone"),
        lab("Estimate", "h* = h_o + α (h_t − h_c): add back the difference between text-only and corrupted-image states"),
        lab("Training-free", "applied at inference; weights unchanged. A defense, so we do not build it; we reuse the size of h_t − h_c per layer as a measurement.")],
        size=14, after=11, name="c method")

    # 5 CMRM PCA and Table 1
    s = new("CMRM, 2 OF 2", "Reading the PCA, and the cost of the fix", "c_pca")
    xr, yb = figure(s, CROPS, "cmrm_f1", 0.6, 1.5, 6.6, 5.3, "CMRM, Fig. 1")
    add_text(s, 7.6, 1.5, 5.1, 2.2, ["PCA squeezes hidden states into two directions of largest variance so clusters show.",
        "Text-only and caption points form one cloud; every input with an image, even blank, sits apart."], size=14, bullet=True, after=7, name="pca read")
    yb = add_table(s, [["LLaVA-1.5-7B", "VLSafe", "Blank", "L-Bench"], ["No intervention", "61.53", "56.87", "79.2"], ["CMRM sample-level", "3.15", "1.24", "77.3"]],
                   7.6, 3.9, [2.2, 1.0, 0.9, 1.0], size=13, pad=0.08, bold_first=True)
    add_text(s, 7.6, yb + 0.2, 5.1, 1.8, ["Unsafe rate falls to near zero with a small capability cost (S-QA 68.0 to 66.1). It shows images shift states, not which property of the image matters."], size=13, color=MUTED, name="c cost")

    # 6 JailBound probe
    s = new("JAILBOUND (NEURIPS 2025), 1 OF 2", "A linear probe: boundary, normal, distance", "j_probe")
    xr, yb = add_fig(s, FIGDIR / "jb_schematic.png", 0.6, 1.5, 6.2, 3.9, "Schematic of a linear probe (ours)")
    add_text(s, 0.6, yb + 0.5, 6.2, 1.0, ["P(unsafe) = sigmoid(w·h + b); boundary w·h + b = 0; normal v = w/|w|; ε = distance to the boundary. One probe per layer."], size=13, name="probe formula")
    xr, yb = figure(s, CROPS, "jb_f7", 7.3, 1.5, 5.4, 3.4, "JailBound, Fig. 7 (appendix): probe training by layer")
    add_text(s, 7.3, yb + 0.25, 5.4, 1.9, ["Deep layers: near-perfect. Layers 0 to 4: poor.", "A probe shows information is readable, not that the model will refuse. We aim the same method at refusal."], size=13, bullet=True, after=6, name="probe read")

    # 7 JailBound Table 3
    s = new("JAILBOUND, 2 OF 2", "Each channel helps, both together help more", "j_t3")
    rows = JB_T3[:5]
    add_text(s, 0.6, 1.45, 7.0, 0.3, ["TABLE 3, ILLEGAL ACTIVITY: WHITE-BOX ASR (%), STRING MATCH"], size=12, color=ACCENT, bold=True)
    bar_chart(s, 0.6, 1.8, 7.4, 4.9, [r[1] for r in rows], [("Llama-3.2-11B", [float(r[2]) for r in rows]), ("Qwen2.5-VL-7B", [float(r[3]) for r in rows]), ("MiniGPT-4", [float(r[4]) for r in rows])],
              horizontal=False, fmt="0", size=11, vmax=115)
    add_text(s, 8.4, 1.5, 4.3, 5.4, [lab("I0 / I1", "clean / perturbed image"), lab("T0 / T1", "original text / text with optimised suffix"),
        lab("ASR", "share of replies with no refusal phrase. Evasive or gibberish replies count, so it is not comparable with judged ASRs."),
        lab("Baseline", "Qwen2.5-VL starts at 2.94%, Llama-3.2 at 51.47%"), lab("Needs", "gradients, so we do not use the attack")], size=13, after=9, name="j text")

    # 8 SafetyReminder Fig 2
    s = new("SAFETYREMINDER (AAAI 2026), 1 OF 2", "Where the safety words fall in a reply", "r_fig2")
    xr, yb = figure(s, CROPS, "sr_f2", 0.6, 1.5, 7.2, 4.9, "SafetyReminder, Fig. 2")
    add_text(s, 8.1, 1.5, 4.6, 5.4, [lab("Safety token", "a refusal or warning word such as “sorry”, “illegal”"), lab("Relative position", "0 at the start of the reply, 1 at the end"),
        lab("Prefilling attack", "forces the reply to start with “Sure, here is…”"),
        lab("Reading", "Qwen2-VL-7B refuses near the start; the attack moves safety words to the end. LLaVA-1.5-7B has them near the end either way.")], size=14, after=11, name="r terms")

    # 9 SafetyReminder tables
    s = new("SAFETYREMINDER, 2 OF 2", "Attack gap between models, and reporting the cost", "r_t1")
    yb = add_table(s, [["ASR %", "Setting", "FigStep", "MMSafety", "SPA-VL", "ImgJP", "GCG", "BAP", "Avg"]] + SR_T1, 0.6, 1.5, [1.9, 1.9, 1.1, 1.2, 1.1, 1.0, 1.0, 1.0, 1.0], size=13, pad=0.06, bold_first=True)
    add_text(s, 0.6, 5.3, 6.0, 1.6, ["FigStep gets 74.9% on LLaVA-1.5-7B but 26.6% on Qwen2-VL-7B: a warning about floor effects for our small Qwen pilot (not yet measured).",
        "Their trained soft prompt (SAPT) is not something we use."], size=13, bullet=True, after=7, name="r read")
    add_text(s, 7.0, 5.3, 5.7, 1.6, ["Utility and over-refusal are reported together: SAPT leaves MM-Vet at 28.0 and 52.6 and benign refusals at 3.7% and 4.1%; BlueSuffix cuts Qwen2-VL to 13.4 with 35.8% refusals.",
        "Their safety-awareness classifier parallels our refusal direction."], size=13, bullet=True, after=7, name="r read2")

    # 10 synthesis
    s = new("SYNTHESIS", "What papers 5 to 8 give us, and what we leave out", "synth")
    yb = add_table(s, [["Paper", "Key evidence", "We take", "We leave out"],
        ["MMJailBench", "Meaning adds 8 to 13 points; OCR 40.6% vs TEXT 50.7%", "TEXT arm; legibility covariate", "Building a benchmark"],
        ["CMRM", "Any image, even blank, shifts the state and the unsafe rate", "Clean-image arm; per-layer shift vs text-only", "The correction (a defense)"],
        ["JailBound", "Safe/unsafe linearly readable; weak in layers 0 to 4", "Layer-wise linear analysis aimed at refusal", "Gradient attack; string-match ASR"],
        ["SafetyReminder", "Qwen2-VL refuses early; prefilling moves it to the end", "First-token and early-token refusal probability", "Trained soft prompt"]],
        0.6, 1.5, [1.9, 4.0, 3.5, 2.7], size=14, pad=0.1, bold_first=True)
    add_text(s, 0.6, yb + 0.3, 12.1, 1.2, ["Three of the four contain a method that trains or uses gradients; only their measurements are borrowed. Still to read: Arditi et al. on refusal directions and the typographic-sweep papers outside these eight."], size=14, color=MUTED, name="synth note")

    OUT.mkdir(parents=True, exist_ok=True)
    prs.save(OUT / f"{NAME}.pptx")
    write_notes(OUT / f"{NAME}_Notes.md", "Papers 5 to 8 in depth: speaker notes",
                "Ten slides: MMJailBench (2), CMRM (2), JailBound (2), SafetyReminder (2), and a synthesis.", titles, N)
    print(f"Saved {(OUT / (NAME + '.pptx')).relative_to(ROOT)} ({n[0]} slides)")


if __name__ == "__main__":
    build()
