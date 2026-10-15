#!/usr/bin/env python3
"""Make Cross_Modal_Conflict_14_Slides_v2.pptx from the user's edited 14-slide deck (which is never modified).

  python3 tools/build_presentation_14_v2.py
What changes in the copy: slides 3-8 rebuilt around exact paper figures/tables (with outlined cells), slides 10/11/13/14
trimmed, a 'question to ask' strip on slides 4-14, accuracy fixes on slides 2 and 13. Slides 1, 9 and 12 keep the user's edits.
"""
import hashlib
import shutil
import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parent))
from deepdive_engine import ACCENT, MUTED, add_table, add_text, extract_crops, figure, lab  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DIR = ROOT / "presentations" / "07_PRESENT_THIS_WEEK_14_SLIDES"
SRC = DIR / "Cross_Modal_Conflict_14_Slides.pptx"
DST = DIR / "Cross_Modal_Conflict_14_Slides_v2.pptx"

CROPS = {
    "figstep_t1": ("figstep", 5, 310, 82, 566, 350), "figstep_f2": ("figstep", 4, 30, 81, 582, 254),
    "vls_f2": ("vlsbench_acl", 3, 298, 205, 540, 376), "vls_t1": ("vlsbench_acl", 3, 298, 60, 540, 135),
    "mmjb_f2": ("mmjailbench", 3, 48, 48, 564, 334), "mmjb_t23": ("mmjailbench", 6, 312, 322, 562, 592),
    "cmrm_f1": ("cmrm", 3, 66, 66, 530, 408), "cmrm_t1_top": ("cmrm", 6, 64, 66, 534, 172),
    "sr_f2": ("safetyreminder_aaai", 2, 50, 50, 298, 216), "sr_t1": ("safetyreminder_aaai", 5, 50, 52, 565, 292),
}
NEW = {k: CROPS[k] for k in ("mmjb_t23", "cmrm_t1_top", "sr_t1")}

ASK = {
    4: "Do we report single-shot ASR, or match FigStep's best-of-5 so our numbers are comparable with theirs?",
    5: "Can we reuse items from VLSBench or MMJailBench as test inputs without that counting as building a benchmark?",
    6: "Prompt framing moves ASR by over 40 points, but we hold the wording fixed. Should we vary framing too?",
    7: "Is a clean-image arm enough to separate “an image is present” from “the image carries an instruction”?",
    8: "How many early tokens should we inspect, and is a 2B model enough if Qwen2-VL-7B resists typographic attacks?",
    9: "If time runs short, which of H1, H2, H3 should we drop first?",
    10: "Is a small hand-written set (6 harmful + 3 benign instructions) acceptable as a controlled stimulus set?",
    11: "Is an LLM judge plus a human audit of about 100 responses enough, or do you want a larger audit?",
    12: "Is a difference-of-means refusal direction acceptable, or should we also fit a probe?",
    13: "OK to download about 4.5 GB of model weights for Gate 0?",
    14: "Is a 5 to 6 week plan realistic, and which threat to validity worries you most?",
}

NOTES = {
 3: """Five papers carry the argument; HADES, SIUO and JailBound were audited but are not presented. If asked: HADES's strongest result needs gradient noise we do not use, SIUO matters only for the judge (harm can arise with no harmful words), JailBound's probes target harmful content rather than refusal and its ASR is a string match.
Three of the five are benchmark papers (VLSBench, MMJailBench, and FigStep's SafeBench). We are not building a benchmark; we take their design lessons, not their data.""",
 4: """What FigStep does (Fig. 2): rewrite a harmful question as the start of a numbered list, print it as an image, and send it with a harmless text prompt asking the model to fill the list in. The harmful words are never in the text.
Table 1 (outlined cells): average ASR 44.80% with the plain question and 82.50% with FigStep, over six open models. ASR is judged by hand, and FigStep was launched five times per question: a question counts as a success if any of the five replies is harmful. The paper does not say the plain-text baseline also got five tries, so the gap may be overstated. Its Fig. 6a shows that with one try ASR is still 82% on LLaVA and about 60% on MiniGPT4 and CogVLM.
What it does not show: a single-try comparison, or anything about image visibility. One image style (black text, white background, size 80, 760 x 760).
Likely question: 'Why not just use FigStep's numbers?' Answer: different judge, best-of-5 protocol, older models; we report single-shot and name the judge.
Ask: single-shot versus best-of-5.""",
 5: """VLSBench does two things. First it shows the problem: a judge labels each pair with image plus text, and again with the text alone. R1 is the share unsafe with both; R2 the share also unsafe from the text alone; R1 minus R2 is harm that needs the image. In four older benchmarks it is 0.27 to 0.91 points (Table 1, outlined), judged by LlamaGuard3-Vision, so the text already carried the harm (Fig. 2 shows examples). Second it builds a fix: a four-step pipeline that writes harmful queries, detoxifies the text so it no longer reveals the harm, generates images, and filters with a model and human checks, giving 2,241 leak-free pairs.
It is a benchmark paper. We are not building a benchmark; we take the design rule that the user text carries no harm cue. R1 and R2 describe datasets, not models. Its 'safety rate' counts warnings as safe, so a high rate can be mostly warnings.
Likely question: 'Why present a benchmark paper if you are not building one?' Answer: it is where a measurement mistake was found, and we avoid it.
Ask: can we reuse benchmark items as inputs.""",
 6: """MMJailBench is a factorial benchmark: 272 harmful intents x 6 prompt framings x 5 visual contexts x 2 carriers (plain text or an image of the text) = 16,320 cases per model, 16 models, judged by GPT-5 (agreement with human experts: QWK 0.97 on harmfulness, 96.3% on detecting misreads, Table 12).
Table 2 (outlined): meaningful visual context adds 8.45 to 12.96 ASR points over no image; blank, noise and nature images add at most 2.19. Table 3: OCR carrier 40.59% ASR against 50.70% for plain text, and still lower after removing misreads (CASR 42.34 vs 50.89). So an instruction rendered in an image is not automatically stronger.
Framing is the biggest factor, more than 40 points between best and worst style. We hold wording fixed, which is a limit.
Likely question: 'Does this contradict FigStep?' Answer: not exactly; FigStep compares to a weak baseline on older models, MMJailBench compares carriers on 16 current models.
Ask: should we vary framing.""",
 7: """CMRM feeds the same harmful question to LLaVA in five forms: text only, text plus the image's caption, a noise image, a blank image, and the original image. Fig. 1 plots their hidden states in 2D (PCA): text-only and caption points form one cloud; everything with an image, even a blank one, sits apart.
Table 1 (outlined, LLaVA-1.5-7B on VLSafe): unsafe rate 1.91% (text only), 5.52% (caption), 56.87% (blank), 57.21% (noise), 61.53% (original). On JailbreakLLMs the gap is much smaller (12.73 text only vs 21.5 to 25.8 with an image), so the 50-point jump is dataset-dependent. 'Unsafe' is judged by Llama-3.1-8B-Instruct and means the model did not refuse.
Its training-free correction is a defense; we do not build it. We keep the blank-image control and the idea of measuring the hidden-state shift per layer against a text-only baseline.
Likely question: 'Is a blank image really a fair control?' Answer: it separates presence from content, but not legibility; that is why we add benign typographic.
Ask: is the clean-image arm enough.""",
 8: """SafetyReminder's observation: a jailbroken VLM often starts to comply and warns only at the end. Fig. 2 plots where safety words fall in a reply (0 start, 1 end). For Qwen2-VL-7B on their harmful data the words cluster near the start, and a prefilling attack ('Sure, here is...') moves them to the end. For LLaVA-1.5-7B they sit near the end either way.
Table 1 (outlined): FigStep ASR with no defense is 74.9% on LLaVA-1.5-7B and 26.6% on Qwen2-VL-7B. This is judged by LlamaGuard3 with greedy decoding, so it differs from FigStep's own 84% (manual, best of 5). Qwen2-VL-7B is robust to typographic prompts only: its ImgJP ASR is 97.4% in the same table, and ADPO's MultiTrust multimodal-jailbreak column shows 54.3%.
We do not use their trained soft prompt. We take the idea that refusal has a position in the reply, so we read the first token and the early tokens.
Likely question: 'Will a 2B Qwen model even be attackable?' Answer: that is Gate 0; if Qwen resists typographic instructions that is itself a result and LLaVA gives contrast.
Ask: how many early tokens, and is 2B enough.""",
 9: "Ask the advisor which hypothesis to drop first if time runs short. H1 visibility, H2 hierarchy, H3 mechanism. Each has a stated way to be wrong; there are no numeric thresholds because we have no data yet.",
 10: "The stimulus set is controlled and small on purpose: 6 harmful and 3 benign instructions, 6 sizes, 270 prompts. It gates a larger run and is not a benchmark.",
 11: "Transcription is the key label: reading the image text aloud is not an attack. Calibrate on the TEXT arm. Each paper's ASR uses a different judge, so we state ours and keep raw responses.",
 12: "All three measurements are inference-only. Attention is correlational, so causal claims come from interventions. The refusal-direction idea of Arditi et al. is cited from memory and still to be read.",
 13: "Built and checked: generator and runner, validated by a dry run. No model has been run. Main risk is a floor effect for typographic prompts on Qwen2-VL. Gate 0 runs the TEXT and benign typographic arms first.",
 14: "Five to six weeks. Lead with judge validity and legibility as threats. No defense is built here; the findings would only inform future alignment methods.",
}


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def clear(slide, keep=4):
    for sh in list(slide.shapes)[keep:]:
        el = sh._element
        rid = None
        if sh.shape_type == 13:
            rid = el.blip_rId
        elif getattr(sh, "has_chart", False) and sh.has_chart:
            rid = el.xpath(".//c:chart/@r:id")[0]
        el.getparent().remove(el)
        if rid:
            slide.part.drop_rel(rid)


def set_para(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def shape_by(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def drop(slide, *names):
    for n in names:
        el = shape_by(slide, n)._element
        el.getparent().remove(el)


def ask(slide, n):
    add_text(slide, 0.6, 6.6, 0.7, 0.3, ["ASK"], size=12, color=ACCENT, bold=True)
    add_text(slide, 1.3, 6.55, 11.4, 0.45, [ASK[n]], size=14, name=f"ask {n}")


def fig(s, key, x, y, w, h, cap, hl=None):
    return figure(s, CROPS, key, x, y, w, h, cap, hl)


tag = lab


def main():
    before = sha(SRC)
    shutil.copy(SRC, DST)
    extract_crops(NEW)
    prs = Presentation(str(DST))
    sl = list(prs.slides)

    # slide 2: accuracy fix, different benchmarks and judges
    set_para(shape_by(sl[1], "TextBox 12").text_frame.paragraphs[1],
             "FigStep reported 82.5% mean ASR on older open VLMs (manual judge, best of 5). Newer Qwen2-VL-7B looks far less affected by "
             "typographic prompts (4.5% in ADPO, 26.6% in SafetyReminder; different judges): when and why the image wins is open.")

    # slide 3: only the five papers we use
    drop(sl[2], "Table 5")
    yb = add_table(sl[2], [["Paper", "What it did", "The number to know", "What we take"],
        ["FigStep", "Harmful request as a numbered-list image + benign text", "ASR 44.8% to 82.5% (best of 5 tries)", "The conflict stimulus; state the protocol"],
        ["VLSBench", "Showed harm leaks through the text; data pipeline with 2,241 leak-free pairs", "R1 minus R2 under 1 point", "User text carries no harm cue; TEXT arm"],
        ["MMJailBench", "Factorial benchmark: intent x framing x visuals x carrier", "Visuals +8 to +13; OCR 40.6% vs TEXT 50.7%", "Always a TEXT arm; report misreads"],
        ["CMRM", "Hidden states for five input variants", "Blank image alone: unsafe 1.91% to 56.87%", "Clean-image arm; per-layer shift"],
        ["SafetyReminder", "Where safety words fall in a reply", "Qwen2-VL: safety words early, late under prefilling", "Read the first and early tokens"]],
        0.6, 1.55, [1.9, 4.4, 3.4, 2.4], size=14, pad=0.1, bold_first=True)
    add_text(sl[2], 0.6, yb + 0.3, 12.1, 0.8, ["Also audited, not presented: HADES, SIUO, JailBound. Three of these five (VLSBench, MMJailBench, FigStep) are benchmark papers; we take their design lessons, not their data."],
             size=14, color=MUTED, name="map note")
    set_para(shape_by(sl[2], "TextBox 2").text_frame.paragraphs[0], "Five papers we draw on")

    # slide 4 FigStep
    s = sl[3]
    clear(s)
    fig(s, "figstep_t1", 0.6, 1.5, 4.6, 4.8, "FigStep, Table 1", hl=[("44.80%", None), ("82.50%", None)])
    xr, yb = fig(s, "figstep_f2", 5.6, 1.5, 7.1, 2.3, "FigStep, Fig. 2")
    add_text(s, 5.6, yb + 0.35, 7.1, 2.3, [
        tag("ASR", "share of harmful questions with a prohibited reply; a success if any of 5 tries is harmful"),
        tag("RESULT", "mean ASR 44.8% (plain text) to 82.5% (FigStep), six open VLMs"),
        tag("CAUTION", "the baseline's number of tries is not stated; with one try ASR is about 60% on MiniGPT4 and CogVLM (Fig. 6a)"),
        tag("WE TAKE", "the conflict stimulus")], size=14, after=7, name="figstep text")

    # slide 5 VLSBench
    s = sl[4]
    clear(s)
    fig(s, "vls_f2", 0.6, 1.5, 5.6, 4.0, "VLSBench, Fig. 2: the red words leak the harm into the text")
    xr, yb = fig(s, "vls_t1", 6.7, 1.5, 6.0, 1.9, "VLSBench, Table 1", hl=[("0.40", None), ("0.36", None), ("0.91", None), ("0.27", None)])
    add_text(s, 6.7, yb + 0.4, 6.0, 2.6, [
        tag("R1 − R2", "harm that needs the image; under 1 point in all four older benchmarks"),
        tag("METHOD", "four-step pipeline: write queries, detoxify the text, generate images, filter; 2,241 leak-free pairs"),
        tag("WE TAKE", "the design rule (no harm cue in the user text), not the data; we are not building a benchmark")],
        size=14, after=8, name="vls text")

    # slide 6 MMJailBench
    s = sl[5]
    clear(s)
    xr, yb = fig(s, "mmjb_f2", 0.6, 1.5, 6.6, 3.7, "MMJailBench, Fig. 2")
    add_text(s, 0.6, yb + 0.3, 6.6, 1.3, ["272 intents x 6 framings x 5 visual contexts x 2 carriers = 16,320 cases per model; 16 models; GPT-5 judge.",
        "We take: always a TEXT arm; report misreads."], size=13, bullet=True, after=5, name="mmjb left text")
    fig(s, "mmjb_t23", 8.2, 1.5, 4.5, 4.6, "MMJailBench, Tables 2 and 3",
        hl=[("12.96", None), ("8.45", None), ("2.19", None), ("50.70", None), ("40.59", None)])

    # slide 7 CMRM
    s = sl[6]
    clear(s)
    fig(s, "cmrm_f1", 0.6, 1.5, 5.2, 4.0, "CMRM, Fig. 1")
    xr, yb = fig(s, "cmrm_t1_top", 6.1, 1.5, 6.6, 1.6, "CMRM, Table 1 (LLaVA-1.5-7B)", hl=[("61.53", None), ("56.87", None), ("57.21", None), ("5.52", None), ("1.91", None)])
    add_text(s, 6.1, yb + 0.4, 6.6, 2.8, [
        tag("VLSAFE", "unsafe rate 1.91% text only, 56.87% with a blank image, 61.53% original"),
        tag("CAUTION", "on JailbreakLLMs the gap is about 10 points; “unsafe” means “did not refuse”"),
        tag("NOT USED", "its correction is a defense"),
        tag("WE TAKE", "a clean-image arm; per-layer shift against a text-only baseline")], size=14, after=7, name="cmrm text")

    # slide 8 SafetyReminder
    s = sl[7]
    clear(s)
    set_para(shape_by(s, "TextBox 2").text_frame.paragraphs[0], "Safety words come early for Qwen2-VL, late under prefilling")
    xr, yb = fig(s, "sr_f2", 0.6, 1.5, 5.0, 3.4, "SafetyReminder, Fig. 2")
    add_text(s, 0.6, yb + 0.3, 5.0, 1.2, ["Safety token: a refusal word. Position: 0 start, 1 end. Prefilling: forces a “Sure, here is…” start."], size=12, color=MUTED, name="sr terms")
    xr, yb = fig(s, "sr_t1", 5.9, 1.5, 6.8, 3.2, "SafetyReminder, Table 1", hl=[("74.9", None), ("26.6", None)])
    add_text(s, 5.9, yb + 0.4, 6.8, 1.9, [
        tag("READING", "FigStep ASR with no defense: 74.9% on LLaVA-1.5-7B, 26.6% on Qwen2-VL-7B (LlamaGuard3, one greedy run)"),
        tag("CAUTION", "Qwen2-VL-7B resists typographic prompts only; ImgJP gets 97.4% in the same table"),
        tag("WE TAKE", "read the first and early tokens; their trained prompt is not used")], size=13, after=7, name="sr text")

    # trims: slide 10 note, slide 11 last bullet, slide 13 pilot-fix block + ADPO wording, slide 14 alignment block
    drop(sl[9], "TextBox 13")
    tf = shape_by(sl[10], "TextBox 6").text_frame
    for p in list(tf.paragraphs)[4:]:
        p._p.getparent().remove(p._p)
    drop(sl[12], "TextBox 7", "TextBox 8")
    set_para(shape_by(sl[12], "TextBox 10").text_frame.paragraphs[0],
             "Typographic jailbreak ASR (ADPO Table 1, MultiTrust): 22.2% on LLaVA-1.5-7B, 4.5% on Qwen2-VL-7B. But Qwen2-VL-7B is not robust overall: 54.3% on the multimodal-jailbreak column.")
    drop(sl[13], "TextBox 8", "TextBox 9")
    set_para(shape_by(sl[13], "TextBox 2").text_frame.paragraphs[0], "Plan and threats to validity")

    for n in range(4, 15):
        ask(sl[n - 1], n)
    for n, txt in NOTES.items():
        sl[n - 1].notes_slide.notes_text_frame.text = txt

    prs.save(str(DST))
    assert sha(SRC) == before, "source deck changed"
    print(f"Saved {DST.relative_to(ROOT)}; source deck unchanged")


if __name__ == "__main__":
    main()
