"""Insert the data explanation and checked Archify diagram, preserving saved edits."""
from pathlib import Path
from copy import deepcopy
import datetime
import hashlib
import re
import shutil
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT/'presentation/VLM_Research_Explained_2026-10-22.pptx'
PNG = ROOT/'presentation/VLM_Data_To_Evidence_Flow.png'
BG, INK, BLUE, MUTED = 'F5F7FB', '14283F', '315EC6', '526579'
records = []

def text(sl, x, y, w, h, value, size=18, bold=False, color=INK):
    box = sl.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(0)
    tf.margin_top = tf.margin_bottom = Pt(0)
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(8)
    records.append(('text', x, y, w, h, value, size, bold, color))
    return box

def rect(sl, x, y, w, h, color='FFFFFF'):
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(color)
    s.line.fill.background()
    records.append(('rect', x, y, w, h, color))

def notes(sl, value):
    # Some user-edited masters omit the notes body placeholder.
    if sl.notes_slide.notes_text_frame is None:
        fresh = Presentation()
        src = fresh.slides.add_slide(fresh.slide_layouts[6]).notes_slide
        for sh in list(sl.notes_slide.shapes):
            sh._element.getparent().remove(sh._element)
        for sh in src.shapes:
            sl.notes_slide.shapes._spTree.insert_element_before(deepcopy(sh._element), 'p:extLst')
    sl.notes_slide.notes_text_frame.text = value

def base(prs, title, sub):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    for sh in list(sl.shapes):
        sh._element.getparent().remove(sh._element)
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    text(sl, 36, 22, 850, 18, 'DATA AND METHOD', 11, True, MUTED)
    text(sl, 36, 51, 890, 43, title, 27, True)
    text(sl, 36, 100, 886, 40, sub, 17, False, MUTED)
    return sl

def preview_pdf(out, pages):
    from reportlab.lib.utils import simpleSplit
    c = canvas.Canvas(str(out), pagesize=(960, 540))
    for page in pages:
        c.setFillColor(HexColor('#'+BG)); c.rect(0, 0, 960, 540, fill=1, stroke=0)
        for r in page:
            if r[0] == 'rect':
                _, x, y, w, h, color = r
                c.setFillColor(HexColor('#'+color)); c.roundRect(x, 540-y-h, w, h, 8, fill=1, stroke=0)
            elif r[0] == 'image':
                _, x, y, w, h = r
                c.drawImage(str(PNG), x, 540-y-h, w, h)
            else:
                _, x, y, w, h, value, size, bold, color = r
                font = 'Helvetica-Bold' if bold else 'Helvetica'
                c.setFont(font, size); c.setFillColor(HexColor('#'+color))
                yy = 540-y-size
                for paragraph in value.split('\n'):
                    for line in simpleSplit(paragraph, font, size, w):
                        c.drawString(x, yy, line); yy -= size*1.15
                    yy -= 8
        c.showPage()
    c.save()

def replace_text(shape, value):
    paras = list(shape.text_frame.paragraphs)
    if paras and paras[0].runs:
        paras[0].runs[0].text = value
        for j, p in enumerate(paras):
            for r in list(p.runs)[1 if j == 0 else 0:]:
                r.text = ''
    else:
        shape.text = value

def main():
    p = Presentation(DECK)
    assert len(p.slides) == 16, 'Expected the saved 16-slide presentation; do not duplicate inserts.'
    digest = hashlib.sha256(DECK.read_bytes()).hexdigest()
    backup = Path('/private/tmp')/('VLM_before_data_flow_'+datetime.datetime.now().strftime('%Y%m%dT%H%M%S')+'.pptx')
    shutil.copy2(DECK, backup)
    old = list(p.slides)
    originals = {s.slide_id: [sh.text for sh in s.shapes if sh.has_text_frame] for s in old}
    # Crop a verified light-theme browser capture; the delivered HTML stays unchanged.
    capture = ROOT/'audit_work/archify_data_flow/VLM_Data_To_Evidence_Flow.visual-check.1440x900.light.png'
    Image.open(capture).crop((64, 107, 1368, 596)).save(PNG)
    pages = []
    s = base(p, 'Where our data comes from — and what we create',
             'We reuse open benchmark questions, then build controlled versions of the inputs.')
    rect(s, 36, 151, 555, 274)
    text(s, 53, 164, 520, 27, 'Existing data used in the completed studies', 20, True)
    rows = [
        ('MM-SafetyBench', '304 questions per Qwen; 152 for LLaVA\nEight selected scenarios in the historical panel.'),
        ('FigStep / SafeBench', '200 questions per model in the historical comparison.'),
        ('Benign controls', 'XSTest and benign twins where those arms ran:\ncheck whether harmless requests are refused too.'),
    ]
    for y, (label, desc) in zip((205, 278, 343), rows):
        text(s, 53, y, 195, 50, label, 17, True, BLUE)
        text(s, 253, y, 320, 65, desc, 15)
    rect(s, 610, 151, 314, 274, 'E8EEF9')
    text(s, 628, 164, 276, 28, 'What our project creates', 20, True)
    text(s, 628, 205, 275, 206,
         'Controlled text/image variants\nInstruction and image controls\nModel-generated responses\nSeparate guard and reviewer labels', 18)
    rect(s, 36, 442, 888, 62, 'E3EBF6')
    text(s, 51, 450, 858, 51,
         'Recent studies reuse eight seen questions. The window study produced 32 responses and 64 views; human review of those views is pending.', 18, True)
    notes(s, 'We did not invent a new independent benchmark. MM-SafetyBench and FigStep supply existing questions and benchmark materials. '
          'We select a documented subset and create controlled input variants and outputs. Historical MM counts are 304 per Qwen and 152 for LLaVA across eight scenarios; FigStep has 200 per model. '
          'XSTest is a benign control, not an attack benchmark. The recent development studies use eight previously seen MM questions, so they are exploratory. '
          'The matched crossover had seven target arms, while the window study reused only four answering conditions. Eight questions times four conditions gives 32 generated responses. '
          'Each longer response is viewed as its first 256 tokens and its full output, giving 64 dependent review units, not 64 independent questions. '
          'These sample sizes belong to separate studies and must not be added as independent evidence. Original benchmark images, replies and row-level annotations remain private. '
          'Evidence: v6_full/results_v6.json; v7_crossover_core8/results_crossover.json; v7_window_core8/results_window.json.')
    pages.append(records.copy()); records.clear()
    f = base(p, 'How one benchmark question becomes evidence',
             'Compare the same questions across conditions; keep every source of labels separate.')
    f.shapes.add_picture(str(PNG), Pt(36), Pt(149), width=Pt(888), height=Pt(333))
    records.append(('image', 36, 149, 888, 333))
    text(f, 40, 488, 880, 34, 'Automated labels are exploratory. Human review can change the interpretation; GPT is an auxiliary annotator.', 15, True)
    notes(f, 'Read the main path from left to right. Start with an existing benchmark question. Create controlled variants: the recent design keeps full-question content fixed and changes plain/steps wording and text/image carrier. '
          'The answering VLM produces the reply; it is different from the judging models. Save each actual response privately. '
          'The same saved reply feeds three separate assessment paths: Qwen3Guard and WildGuard, the human reviewer, and an auxiliary GPT-6.1-Sol reviewer. '
          'Human and auxiliary LLM review packets hide guard labels and condition codes where possible, although wording and carrier may remain visible. Human and GPT labels are never merged into a second human reference. '
          'The original 48-response human review received LLM assistance on some examples; disclose that limitation. The new 64-unit human window review is pending. '
          'Both guards judge refusal and harmful assistance. A separate historical non-answer heuristic is imperfect; observable comprehension is reviewed under the fixed rubric and can be uncertain. '
          'Finally compare paired outcomes for the same questions. Report refusal, harmful assistance, comprehension evidence, uncertainty, missing labels and truncation separately. '
          'The diagram explains the intended end-to-end flow; it does not imply every human-review stage is complete. Current conclusions remain exploratory. '
          'There is no model training and no new benchmark invented. The saved original responses and labels remain unchanged. Archify dataflow diagram: presentation/VLM_Data_To_Evidence_Flow.html.')
    pages.append(records.copy()); records.clear()
    desired = old[:1]+[s, f]+old[1:]
    order = [sl.slide_id for sl in desired]
    ids = p.slides._sldIdLst; byid = {x.id:x for x in list(ids)}
    for x in list(ids): ids.remove(x)
    for ident in order: ids.append(byid[ident])
    mapping = next(sl for sl in old if any(sh.has_text_frame and 'From prior ideas' in sh.text for sh in sl.shapes))
    def renumber_ref(match):
        return match.group(1)+'/'.join(str(int(n)+2) for n in match.group(2).split('/'))+':'
    for sh in mapping.shapes:
        if sh.has_text_frame:
            replace_text(sh, re.sub(r'(Slides? )(\d+(?:/\d+)*):', renumber_ref, sh.text))
    for i, sl in enumerate(desired, 1):
        footer = False
        for sh in sl.shapes:
            if sh.has_text_frame and sh.top/Pt(1)>500 and re.fullmatch(r'\s*\d{1,2}\s*/\s*\d{1,2}\s*', sh.text):
                replace_text(sh, f'{i:02d} / 18'); footer=True
        if not footer: text(sl, 873, 520, 55, 16, f'{i:02d} / 18', 10, color=MUTED)
    for sl in old:
        before = originals[sl.slide_id]
        after = [sh.text for sh in sl.shapes if sh.has_text_frame]
        for a, b in zip(before, after):
            assert a == b or re.fullmatch(r'\s*\d+\s*/\s*\d+\s*', a) or sl is mapping
        assert len(before) == len(after)
    out = Path('/private/tmp/VLM_data_flow_verified.pptx')
    p.save(out)
    check = Presentation(out); assert len(check.slides) == 18
    assert 'What our project creates' in '\n'.join(sh.text for sh in check.slides[1].shapes if sh.has_text_frame)
    assert check.slides[2].notes_slide.notes_text_frame is not None
    assert hashlib.sha256(DECK.read_bytes()).hexdigest() == digest, 'Deck changed during insertion.'
    shutil.copy2(out, DECK)
    preview_pdf(ROOT/'audit_work/archify_data_flow/new_slides_preview.pdf', pages)
    print('Updated:', DECK, '\nSlides: 18. Added data slide 2 and flow slide 3.\nPreserved original slide text except cross-references/page numbers.\nBackup:', backup)

if __name__ == '__main__': main()
