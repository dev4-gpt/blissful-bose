#!/usr/bin/env python3
"""Build the pilot-study deck: text vs vision safety coupling in open VLMs.

All numbers below come from bundle (1)/results.json (v1 judge) and from the raw-response audit of
2026-10-14. v2 numbers are NOT included; slide 11 only lists what v2 changes.
Usage: python3 tools/build_modality_safety_deck.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "bundle (1)"
OUT = ROOT / "presentation" / "Modality_Safety_Pilot_Deck.pptx"

INK, SLATE, AMBER, TEAL, RED = "0F1B2D", "24364F", "F2A33A", "2A9D8F", "D1495B"
GRAY, LIGHT, WHITE, MID = "8A97A8", "F2F5F8", "FFFFFF", "5B6B80"
HEAD, BODY = "Cambria", "Calibri"


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
_n = [0]


def new_slide(dark=False, notes=""):
    s = prs.slides.add_slide(BLANK)
    _n[0] += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(INK if dark else WHITE)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    if _n[0] > 1:
        text(s, str(_n[0]), 12.3, 6.95, 0.7, 0.35, size=12, color=GRAY, align=PP_ALIGN.RIGHT)
    return s


def text(s, t, x, y, w, h, size=18, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, line=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = t if isinstance(t, list) else [t]
    for i, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        if i:
            p.space_before = Pt(6)
        r = p.add_run()
        r.text = ptxt
        r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
        r.font.color.rgb, r.font.name = rgb(color), font
    return tb


def box(s, x, y, w, h, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, radius=0.08):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(1.25)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def chip(s, label, fill=AMBER, color=INK, x=0.6, y=0.38):
    w = 0.14 * len(label) + 0.5
    b = box(s, x, y, w, 0.34, fill=fill, radius=0.5)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.name = Pt(13), True, BODY
    r.font.color.rgb = rgb(color)


def title(s, t, color=INK):
    text(s, t, 0.6, 0.85, 12.1, 1.0, size=32, bold=True, color=color, font=HEAD, anchor=MSO_ANCHOR.TOP)


def stat(s, x, y, w, big, label, big_color=AMBER, label_color=INK, bigsize=48):
    text(s, big, x, y, w, 0.9, size=bigsize, bold=True, color=big_color, font=HEAD)
    text(s, label, x, y + 0.95, w, 0.9, size=16, color=label_color)


def style_chart(chart, colors, fmt="0.00", legend=True, maxv=1.0):
    chart.has_title = False
    chart.font.size = Pt(14)
    chart.font.name = BODY
    chart.font.color.rgb = rgb(INK)
    va = chart.value_axis
    va.maximum_scale, va.minimum_scale = maxv, 0
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("DDE3EA")
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(12)
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.0", False
    va.tick_labels.font.color.rgb = rgb(MID)
    ca = chart.category_axis
    ca.format.line.color.rgb = rgb(GRAY)
    ca.tick_labels.font.size = Pt(14)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(14)
    plot = chart.plots[0]
    plot.gap_width, plot.overlap = 60, -10
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold = Pt(13), True
    for ser, col in zip(plot.series, colors):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(col)


# ---------------------------------------------------------------- 1 title
s = new_slide(dark=True, notes=(
    "Open with the question, not the method. Today's message in one line: we ran a first pilot, an audit of our own "
    "pipeline caught a serious measurement problem, and we now know exactly what to fix. Be clear these are PILOT numbers."))
chip(s, "PILOT RESULTS, NOT FINAL", fill=AMBER, y=0.9)
text(s, "Do text and vision safety trade off in open vision-language models?", 0.6, 1.7, 11.6, 2.4,
     size=44, bold=True, color=WHITE, font=HEAD)
text(s, "A pilot with frozen open models, and what auditing our own results taught us", 0.6, 4.2, 11, 0.9,
     size=22, color="C9D3E0")
text(s, "Aryaman Singh Dev  |  Advisor: Prof. Thao Minh Le  |  Penn State  |  October 15, 2026", 0.6, 6.3, 12, 0.5,
     size=16, color=GRAY)

# ---------------------------------------------------------------- 2 question
s = new_slide(notes=(
    "Three competing readings of the advisor's idea. The see-saw says safety is a fixed budget moved between modalities. "
    "Most of the literature points to shared circuitry in the language backbone. A third possibility is independence. "
    "Our design can tell these apart: same request, text versus image, in the same model."))
chip(s, "THE QUESTION")
title(s, "Three ways text and vision safety could relate")
cards = [("See-saw", "Vision safety drops by Z, text safety rises by Z. Safety is a budget moved between channels.", AMBER),
         ("Shared circuitry", "Both channels use one refusal mechanism in the language backbone, so they move together.", TEAL),
         ("Independent", "The channels do not predict each other.", GRAY)]
for i, (h, b, col) in enumerate(cards):
    x = 0.6 + i * 4.1
    box(s, x, 2.15, 3.85, 3.0, fill=LIGHT)
    box(s, x + 0.3, 2.45, 0.5, 0.5, fill=col, shape=MSO_SHAPE.OVAL)
    text(s, str(i + 1), x + 0.3, 2.45, 0.5, 0.5, size=20, bold=True, color=INK, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, h, x + 0.3, 3.15, 3.3, 0.5, size=24, bold=True, font=HEAD)
    text(s, b, x + 0.3, 3.75, 3.3, 1.3, size=20, color=SLATE)
text(s, "Test: the same harmful request, once as text and once as an image, in frozen open models. No training.",
     0.6, 5.65, 12.1, 0.8, size=20, bold=True, color=INK)

# ---------------------------------------------------------------- 3 literature
s = new_slide(notes=(
    "Read from the abstracts and two papers in full; I verified one key paper myself (MARS). Everything on this slide is "
    "'what is already known'. The important point is the last row: refusal-direction transfer is already published, so "
    "steering is a replication for us, not our contribution."))
chip(s, "WHAT IS KNOWN")
title(s, "The literature already covers most of this")
rows = [("Adding vision erodes safety", "Learning to See but Forgetting to Follow; VLGuard (ICML 2024)"),
        ("Images sit in a different representation region", "Liu et al. (ACL Findings 2025); TGA (ICLR 2025)"),
        ("Typographic images slip past text safety", "FigStep (AAAI 2025); MM-SafetyBench; Text2VLM"),
        ("Refusal directions transfer across modalities", "MARS (arXiv 2606.31876, Jun 2026); Arditi et al. (2024)")]
for i, (h, c) in enumerate(rows):
    y = 2.05 + i * 1.15
    box(s, 0.6, y, 12.1, 1.0, fill=LIGHT)
    box(s, 0.85, y + 0.25, 0.5, 0.5, fill=TEAL if i < 3 else AMBER, shape=MSO_SHAPE.OVAL)
    text(s, str(i + 1), 0.85, y + 0.25, 0.5, 0.5, size=18, bold=True, color=INK, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, h, 1.65, y + 0.14, 5.4, 0.8, size=19, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, c, 7.2, y + 0.14, 5.3, 0.8, size=18, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
text(s, "Steering transfer is already published, so it is a replication for us.", 0.6, 6.65, 11, 0.4, size=16,
     bold=True, color=INK)

# ---------------------------------------------------------------- 4 design
s = new_slide(notes=(
    "One harmful request from MM-SafetyBench appears in four forms. A and B differ only by model (LLM backbone versus VLM). "
    "C moves the harmful phrase into an image. D adds a blank image to control for 'an image is present'. "
    "Judge is a small open model. Greedy decoding. Everything runs on a free Colab T4."))
chip(s, "DESIGN")
title(s, "One request, four channels")
conds = [("A", "LLM backbone", "Harmful question as text", GRAY),
         ("B", "VLM", "Same question as text, no image", TEAL),
         ("C", "VLM", "Harmful phrase typed in an image + generic prompt", AMBER),
         ("D", "VLM", "Same question + blank image", MID)]
for i, (k, m, d, col) in enumerate(conds):
    x = 0.6 + i * 3.08
    box(s, x, 2.1, 2.9, 2.75, fill=LIGHT)
    box(s, x + 0.25, 2.35, 0.6, 0.6, fill=col, shape=MSO_SHAPE.OVAL)
    text(s, k, x + 0.25, 2.35, 0.6, 0.6, size=22, bold=True, color=WHITE if col in (MID, GRAY, TEAL) else INK,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, m, x + 0.25, 3.1, 2.4, 0.45, size=20, bold=True, font=HEAD)
    text(s, d, x + 0.25, 3.6, 2.45, 1.2, size=20, color=SLATE)
text(s, ["MM-SafetyBench: 8 harmful scenarios x 30 items = 240 per condition",
         "Models: Qwen2-VL-2B and Qwen2.5-VL-3B, with their text backbones",
         "Judge: open Qwen2.5-3B; greedy decoding; free Colab T4"], 0.6, 5.2, 12.1, 1.5, size=22, color=SLATE)

# ---------------------------------------------------------------- 5 result 1
s = new_slide(notes=(
    "As measured with the first judge. Direction: the image channel refuses less than text in both models. "
    "The size of that gap is UNDERSTATED here because the judge mislabelled image-condition compliance as refusal; "
    "slide 8 explains. The blank-image bar for the 2B model is higher than plain text; the audit shows that is the model "
    "switching to short canned refusals, not better judgement."))
chip(s, "PILOT  |  JUDGE V1", fill=AMBER)
title(s, "The image channel refuses less than text")
cd = CategoryChartData()
cd.categories = ["A  LLM text", "B  VLM text", "D  VLM + blank image", "C  VLM image"]
cd.add_series("Qwen2-VL-2B", (0.796, 0.721, 0.842, 0.521))
cd.add_series("Qwen2.5-VL-3B", (0.812, 0.821, 0.821, 0.613))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.9), Inches(8.2), Inches(4.9), cd)
style_chart(gf.chart, [AMBER, TEAL])
text(s, "Refusal rate (higher = safer), n = 240 per bar", 0.6, 6.8, 8.2, 0.35, size=13, color=MID)
box(s, 9.2, 2.1, 3.5, 4.4, fill=LIGHT)
text(s, "Direction holds", 9.45, 2.3, 3.0, 0.5, size=22, bold=True, font=HEAD)
text(s, ["Image refusal is below text refusal in both models.",
         "The size is understated: the judge mislabelled image-condition answers (slide 8)."],
     9.45, 2.95, 3.05, 3.4, size=20, color=SLATE)

# ---------------------------------------------------------------- 6 coupling
s = new_slide(notes=(
    "Each dot is a scenario: text refusal on one axis, image refusal on the other. They rise together, which argues against a "
    "see-saw. Caution: scenarios differ in difficulty, so topic alone can create this pattern. The v2 run adds a "
    "within-scenario correlation to separate the two. Image numbers use the v1 judge."))
chip(s, "PILOT  |  JUDGE V1", fill=AMBER)
title(s, "Text and image refusal rise and fall together")
pic = BUNDLE / "fig2_coupling.png"
if pic.exists():
    s.shapes.add_picture(str(pic), Inches(0.6), Inches(1.9), height=Inches(4.7))
stat(s, 6.8, 2.0, 2.8, "0.89 / 0.73", "Scenario-level Spearman, 2B / 3B (8 scenarios)", big_color=TEAL, bigsize=36)
stat(s, 6.8, 3.9, 2.8, "0.50 / 0.54", "Item-level correlation (phi), 2B / 3B", big_color=TEAL, bigsize=36)
text(s, "Not a see-saw. But topic difficulty alone could produce this; v2 tests it within scenario.", 9.9, 2.1,
     2.9, 3.6, size=22, color=SLATE)

# ---------------------------------------------------------------- 7 steering
s = new_slide(notes=(
    "Steering adds a direction learned from text refusals to the model's internal activations. In the pilot it raised refusal on "
    "harmful images too, but it also made the model refuse harmless questions. This is a replication of published work, "
    "and v1 had no random-direction control, so we will not claim it as a finding yet."))
chip(s, "PILOT  |  JUDGE V1", fill=AMBER)
title(s, "A text-derived direction moves images too, at a cost")
pic = BUNDLE / "fig3_steering.png"
if pic.exists():
    s.shapes.add_picture(str(pic), Inches(0.6), Inches(1.9), height=Inches(4.7))
text(s, ["At strength 0.5:", "harmful text 0.65 to 0.91", "harmful image 0.50 to 0.86", "harmless prompts 0.00 to 0.50"],
     7.9, 2.1, 4.8, 2.5, size=20, color=INK)
box(s, 7.9, 4.5, 4.8, 1.9, fill=LIGHT)
text(s, "Already published (MARS). No random-direction control in this run, so it is a replication, not a result.",
     8.15, 4.7, 4.3, 1.6, size=20, color=SLATE)

# ---------------------------------------------------------------- 8 audit
s = new_slide(dark=True, notes=(
    "This is the most important slide. An independent audit agent read the raw responses and told us the judge was failing on the "
    "image condition. I re-checked it myself: in the Qwen2-VL-2B image condition, 90 of 125 'refusals' contained no refusal "
    "phrase at all; they were step-by-step compliance. For the 3B model it was 40 of 147. In conditions A, B and D the judge and a "
    "simple phrase check agree within a few points. The phrase check is a proxy, not ground truth; v2 adds a human-labelled "
    "sample. Net effect: the true image-channel gap is larger than we reported."))
chip(s, "AUDIT  |  CHECKED ON RAW RESPONSES", fill=RED, color=WHITE)
title(s, "Our judge failed on the image condition", color=WHITE)
stat(s, 0.6, 2.0, 3.6, "90 of 125", "Qwen2-VL-2B image-condition 'refusals' with no refusal phrase", big_color=AMBER,
     label_color="C9D3E0", bigsize=40)
stat(s, 0.6, 4.0, 3.6, "40 of 147", "Same check, Qwen2.5-VL-3B", big_color=AMBER, label_color="C9D3E0", bigsize=40)
cd = CategoryChartData()
cd.categories = ["Qwen2-VL-2B", "Qwen2.5-VL-3B"]
cd.add_series("Judge says refusal", (0.521, 0.613))
cd.add_series("Refusal phrase present", (0.146, 0.446))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4.7), Inches(1.9), Inches(5.0), Inches(4.4), cd)
style_chart(gf.chart, [AMBER, TEAL])
c = gf.chart
c.font.color.rgb = rgb(WHITE)
c.value_axis.tick_labels.font.color.rgb = rgb("C9D3E0")
c.value_axis.major_gridlines.format.line.color.rgb = rgb("3A4D66")
c.legend.font.color.rgb = rgb(WHITE)
c.plots[0].data_labels.font.color.rgb = rgb(WHITE)
text(s, "Image condition (C)", 4.7, 6.35, 5.0, 0.4, size=13, color=GRAY)
text(s, ["Judge counted step-by-step compliance as refusal.",
         "In conditions A, B, D the judge and the phrase check agree within a few points.",
         "True image-channel gap is larger than reported."], 10.0, 2.1, 2.8, 4.3, size=18, color="C9D3E0")

# ---------------------------------------------------------------- 9 claims
s = new_slide(notes=(
    "A claims ladder: only say what the data supports. The blank-image result: 187 of 240 answers were the same short canned "
    "refusal, so it is a style switch, not better safety. Do not claim 'visual context improves safety'."))
chip(s, "WHAT WE MAY CLAIM")
title(s, "What the pilot supports, and what it does not")
box(s, 0.6, 2.0, 5.95, 4.6, fill=LIGHT)
text(s, "Supported (direction)", 0.9, 2.2, 5.4, 0.5, size=22, bold=True, color=TEAL, font=HEAD)
text(s, ["The image channel refuses less than text, in both models",
         "Text and image refusal are positively coupled",
         "A blank image shifts the 2B model into canned refusals"], 0.9, 2.9, 5.35, 3.5, size=22, color=SLATE)
box(s, 6.75, 2.0, 5.95, 4.6, fill=LIGHT)
text(s, "Not yet supported", 7.05, 2.2, 5.4, 0.5, size=22, bold=True, color=RED, font=HEAD)
text(s, ["The size of the image gap (judge was wrong)",
         "Modality versus prompt wording: image prompt differs from text prompt",
         "A causal shared circuit, or steering transfer (no random-direction control)"],
     7.05, 2.9, 5.4, 3.5, size=22, color=SLATE)

# ---------------------------------------------------------------- 10 novelty
s = new_slide(notes=(
    "Novelty check by a literature agent, spot-verified on MARS. The search was not exhaustive, so we do not claim 'first'. "
    "Our defensible contribution is a careful measurement study: how coupled the channels are, how big the gap is once measured "
    "correctly, and what steering costs in over-refusal."))
chip(s, "NOVELTY CHECK")
title(s, "Where we sit against published work")
box(s, 0.6, 2.0, 5.95, 4.6, fill=LIGHT)
text(s, "Already published", 0.9, 2.2, 5.4, 0.5, size=22, bold=True, font=HEAD)
text(s, ["Refusal-direction transfer to images (MARS)", "Vision erodes safety; typographic gap",
         "Blank and noise image controls on LLaVA"], 0.9, 2.9, 5.35, 3.5, size=22, color=SLATE)
box(s, 6.75, 2.0, 5.95, 4.6, fill=LIGHT)
text(s, "Possibly open", 7.05, 2.2, 5.4, 0.5, size=22, bold=True, color=TEAL, font=HEAD)
text(s, ["Per-item coupling of text and image refusal", "Typographic channel inside direction transfer",
         "Blank-image behaviour on Qwen-VL models", "Search was not exhaustive: no 'first' claim yet"],
     7.05, 2.9, 5.4, 3.5, size=22, color=SLATE)

# ---------------------------------------------------------------- 11 v2
s = new_slide(notes=(
    "v2 is running on a Colab T4 now. Do not quote any v2 number until it finishes and the hand-labelled audit confirms the judge. "
    "If v2 finishes before the talk, add its numbers here."))
chip(s, "V2  |  RUNNING NOW, NO RESULTS YET", fill=TEAL, color=WHITE)
title(s, "v2 fixes every flaw the audit found")
fixes = [("Fix the judge", "Stricter rubric, plus a refusal-phrase gate"),
         ("Prompt-matched control", "Same prompt, phrase as text (condition E)"),
         ("Harmless twins", "48 benign prompts measure over-refusal"),
         ("Honest statistics", "Random ids, scenario-clustered CIs, within-scenario coupling"),
         ("Steering control", "Random direction of equal size"),
         ("Human labels", "About 240 hand-labelled answers, Cohen's kappa")]
for i, (h, b) in enumerate(fixes):
    x, y = 0.6 + (i % 3) * 4.1, 2.0 + (i // 3) * 2.45
    box(s, x, y, 3.85, 2.2, fill=LIGHT)
    box(s, x + 0.25, y + 0.25, 0.5, 0.5, fill=AMBER, shape=MSO_SHAPE.OVAL)
    text(s, str(i + 1), x + 0.25, y + 0.25, 0.5, 0.5, size=18, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, h, x + 0.25, y + 0.9, 3.4, 0.45, size=19, bold=True, font=HEAD)
    text(s, b, x + 0.25, y + 1.35, 3.4, 0.8, size=18, color=SLATE)

# ---------------------------------------------------------------- 12 roadmap
s = new_slide(notes=(
    "Deadlines are from the venue agent, which opened official pages; confirm the ARR January date when it is posted. "
    "ACL 2027 via ARR on January 4, 2027 is the target. The AAAI-27 workshops on November 20 are non-archival and none is "
    "dedicated to VLM safety. TMLR is rolling."))
chip(s, "PATH TO A PAPER")
title(s, "Target: ACL 2027 short paper, due January 4")
steps = [("Oct 15-18", "Hand-label answers, fix judge, add controls"),
         ("Oct 19-Nov 5", "Rerun both pairs, steering controls, first draft"),
         ("Nov 20", "Optional: AAAI-27 workshop, 4 pages"),
         ("Dec 20", "n=100 per scenario, third model family, full draft"),
         ("Jan 4", "Submit via ACL Rolling Review")]
for i, (d, t) in enumerate(steps):
    x = 0.6 + i * 2.45
    box(s, x, 2.3, 2.3, 3.4, fill=LIGHT)
    text(s, d, x + 0.2, 2.5, 1.9, 0.5, size=20, bold=True, color=TEAL if i != 4 else RED, font=HEAD)
    text(s, t, x + 0.2, 3.15, 1.95, 2.4, size=18, color=SLATE)
text(s, "Fallbacks: TMLR (rolling, reviews correctness); workshops listed Feb 2027.", 0.6, 6.1, 12, 0.5, size=17,
     color=INK, bold=True)

# ---------------------------------------------------------------- 13 asks
s = new_slide(dark=True, notes=(
    "Close with the decisions you need from the advisor. Frame the claim around coupling and the typographic channel, "
    "not steering novelty."))
chip(s, "DECISIONS", fill=AMBER)
title(s, "What we need from you", color=WHITE)
asks = ["Claim framing: channel coupling and the typographic channel, not steering novelty",
        "Target venue: ACL 2027 short paper, with an optional November workshop version",
        "Scope: does no-training activation steering fit the 'no heavy training' rule?",
        "Process: authorship order, and an arXiv endorsement for a first-time submitter"]
for i, a in enumerate(asks):
    y = 2.1 + i * 1.12
    box(s, 0.6, y, 0.6, 0.6, fill=AMBER, shape=MSO_SHAPE.OVAL)
    text(s, str(i + 1), 0.6, y, 0.6, 0.6, size=22, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, a, 1.5, y - 0.02, 11.2, 0.9, size=24, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print("saved", OUT)
