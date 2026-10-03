from hashlib import sha256
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


PPTX = Path(__file__).resolve().parents[1] / "VLM_Safety_Alignment_10_Paper_Proposal.pptx"
INK = RGBColor(25, 25, 25)
MID = RGBColor(75, 75, 75)
LIGHT = RGBColor(220, 220, 220)
PALE = RGBColor(245, 245, 245)


def fingerprint(slide):
    result = []
    for shape in slide.shapes:
        value = [shape.shape_type, shape.left, shape.top, shape.width, shape.height]
        if getattr(shape, "has_text_frame", False):
            value.append(shape.text)
        if shape.shape_type == 13:
            value.append(sha256(shape.image.blob).hexdigest())
        result.append(tuple(value))
    return tuple(result)


def textbox(slide, x, y, w, h, text, size=12, bold=False, color=INK,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font="Aptos"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(.035)
    tf.margin_right = Inches(.035)
    tf.margin_top = Inches(.025)
    tf.margin_bottom = Inches(.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def line(slide, x1, y1, x2, y2, color=INK, width=1):
    shape = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def slide_header(slide, section, title, citation, page):
    textbox(slide, .62, .22, 5.8, .25, section.upper(), 10, True, MID)
    textbox(slide, .62, .53, 12.1, .56, title, 24, True)
    textbox(slide, .64, 1.12, 12.0, .24, citation, 10, False, MID)
    line(slide, .62, 1.48, 12.72, 1.48, INK, 1.1)
    textbox(slide, .62, 7.08, 11.2, .18, "VLM Safety Alignment | Focused literature review", 8.5, False, MID)
    textbox(slide, 12.18, 7.06, .5, .2, f"{page:02d}", 9, False, MID, PP_ALIGN.RIGHT)


def paper_block(slide, x, title, body, side_label):
    textbox(slide, x, 1.72, 5.72, .48, title, 15, True)
    line(slide, x, 2.24, x + 5.72, 2.24, LIGHT, .9)
    textbox(slide, x, 2.32, 5.72, 3.78, body, 12.2)
    textbox(slide, x, 6.27, 5.72, .22, side_label.upper(), 9, True, MID)


def set_notes(prs, slide, text):
    notes = slide.notes_slide
    if notes.notes_text_frame is None:
        template = next(
            (existing.notes_slide for existing in prs.slides
             if existing is not slide and existing.notes_slide.notes_text_frame is not None),
            None,
        )
        if template is None:
            raise RuntimeError("No existing speaker-notes layout is available.")
        body = next(p for p in template.placeholders if p.placeholder_format.idx == 3)
        notes._element.spTree.insert(2, deepcopy(body._element))
    slide.notes_slide.notes_text_frame.text = text


def add_review_slide(prs, data, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "Literature Review | Paper-by-Paper Evidence", data["title"], data["refs"], page)
    paper_block(slide, .66, data["left_title"], data["left_body"], data["left_role"])
    line(slide, 6.66, 1.72, 6.66, 6.55, LIGHT, .8)
    paper_block(slide, 6.94, data["right_title"], data["right_body"], data["right_role"])
    set_notes(prs, slide, data["notes"])
    return slide


def add_attack_coverage_slide(prs, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "Attack Evaluation", "A representative suite, not every published VLM attack",
                 "Attack procedures: FigStep [1], HADES [11], PolyJailbreak [4] | Attack corpus: JailBreakV-28K [12] | Compositional tests: SIUO [13]",
                 page)
    columns = [(.68, 2.22, "FAMILY / ACCESS"), (3.02, 5.0, "WHAT IT TESTS"), (8.18, 4.45, "PLACE IN OUR PROTOCOL")]
    for x, w, header in columns:
        textbox(slide, x, 1.7, w, .3, header, 10, True, MID)
    line(slide, .68, 2.04, 12.62, 2.04, INK, 1)
    rows = [
        ("Visual text / OCR\nBlack-box", "FigStep turns an unsafe instruction into image typography; vary rendering while holding intent fixed.", "CORE fixed tier; text-only matched control + benign OCR controls."),
        ("Generated / optimized image\nBlack-box or white-box", "HADES separates typography, black-box optimized images (+Opt), and gradient-based adversarial images (+Adv). These are different access tiers.", "CORE +Opt if artifacts are reproducible; +Adv is a separate optional white-box result."),
        ("Cross-modal composition\nBenchmark cases", "SIUO and HoliSafe cover cases where modality-wise safe inputs can yield an unsafe joint interpretation; not attack algorithms.", "CORE strata; test each modality alone and jointly, with safe near-neighbors."),
        ("Adaptive cross-modal search\nBlack-box", "PolyJailbreak iteratively adapts text/image strategies from target feedback; report cost and success by query count.", "CORE adaptive tier under a shared target-call budget; keep separate from static ASR."),
        ("Transferred / diverse attack corpus", "JailBreakV-28K is a benchmark collection spanning transferred text and multimodal attacks; the corpus is not one attack method.", "Use a deduplicated, held-out slice as a stress test; disclose selection and overlap."),
    ]
    y = 2.14
    row_h = .77
    for idx, row in enumerate(rows):
        textbox(slide, .68, y + .04, 2.22, row_h - .08, row[0], 11.2, True)
        textbox(slide, 3.02, y + .04, 5.0, row_h - .08, row[1], 10.8)
        textbox(slide, 8.18, y + .04, 4.45, row_h - .08, row[2], 10.8)
        y += row_h
        line(slide, .68, y, 12.62, y, LIGHT, .7)
    textbox(slide, .72, 6.18, 11.8, .46,
            "Out of initial scope: hidden/steganographic image instructions, structured flowchart attacks, multilingual and multi-turn variants. Add as preregistered extensions; do not claim exhaustive coverage.",
            10.2, True)
    set_notes(prs, slide, (
        "There is no finite slide that can cover every VLM attack in the literature, and new attack variants continue to appear. "
        "This is a mechanism-based coverage plan, not an exhaustive claim. The initial core covers visual OCR/typography (FigStep), "
        "image-generation/optimization conditions (HADES +Opt; HADES +Adv only as a separately approved white-box tier), "
        "cross-modal composition (SIUO/HoliSafe benchmark strata), and adaptive black-box search (PolyJailbreak). JailBreakV-28K is "
        "a collection/benchmark of diverse jailbreak instances, not a single procedure; sample it without overlap and describe exactly "
        "which subset is run. Keep image removal, masking, decoys, and benign near-neighbors as attribution/calibration controls, not attacks. "
        "The slide explicitly flags omitted attack families: hidden or steganographic visual instructions, structured graphic/flowchart prompts, "
        "multilingual and multi-turn variants, and visual prompt injection in external content. They are reasonable follow-on tests after the "
        "core run, but require their own source-intent split, threat model, budget, and primary-paper audit. Never describe the current suite as "
        "testing every attack or as a complete leaderboard."
    ))
    return slide


def main():
    prs = Presentation(PPTX)
    if len(prs.slides) != 18:
        raise RuntimeError(f"Expected the saved 18-slide deck, found {len(prs.slides)}; no changes made.")
    before_slide2 = fingerprint(prs.slides[1])
    before_slide3 = fingerprint(prs.slides[2])
    if not any(s.shape_type == 13 for s in prs.slides[1].shapes) or len(prs.slides[2].shapes) != 0:
        raise RuntimeError("Slide 2 has no saved image content or slide 3 is not blank; no changes made.")
    if not any(getattr(s, "has_text_frame", False) and s.text.strip() == "Ten papers, four jobs in the proposed study" for s in prs.slides[5].shapes):
        raise RuntimeError("Expected literature map is not on slide 6; no changes made.")

    review_slides = [
        {
            "title": "Attacks: from fixed visual prompts to adaptive search",
            "refs": "[1] FigStep, §§5–6 / Table 1 / Fig. 2 | [4] PolyJailbreak, method / Table VI / Fig. 6",
            "left_title": "FigStep | fixed typographic jailbreak [1]",
            "left_body": "CONTRIBUTION\nBlack-box attack: paraphrase a prohibited request, render it as numbered text in an image, then pair it with neutral incitement. SafeBench: 500 reviewed questions; SafeBench-Tiny: 50.\n\nEVALUATION + RESULT\nSix open VLMs; manual safety assessment. Table 1 reports mean ASR 82.50% for FigStep vs. 44.80% for the vanilla text query. Each item was tried five times; any successful attempt counted.\n\nPROJECT MAPPING\nFixed OCR/visual-text tier; add matched text-only prompts, benign OCR controls, held-out renderings, and image ablations. It is an attack benchmark, not an alignment method.",
            "left_role": "Pipeline: fixed attack + text/image attribution control",
            "right_title": "PolyJailbreak | adaptive cross-modal attack [4]",
            "right_body": "CONTRIBUTION\nBlack-box attacker composes cross-modal strategy primitives and iteratively adapts using target feedback. The paper evaluates eight open and closed MLLMs with up to 15 optimization steps.\n\nEVALUATION + RESULT\nTable VI reports 83.34% average ASR and 3.976/5 harmfulness score, judged with GPT-4o. These are separate outcomes; its optimization steps are not a universal total-call budget.\n\nPROJECT MAPPING\nAdaptive red-team tier; equalize target-call budgets, log auxiliary calls/cost, and plot success vs. queries. Keep separate from fixed-attack ASR.",
            "right_role": "Pipeline: adaptive, separately budgeted red-team tier",
            "notes": "[1] FigStep: source sections §§5.2 and 6.1–6.3; Figure 2 is the actual three-step method (paraphrase, typography, incitement), Figure 1 is conceptual. Table 1 is the correct source for the 82.50% versus 44.80% mean comparison. SafeBench has 500 items across ten topics; SafeBench-Tiny has 50; each item was attempted five times and one success counted. The reported 82.50/44.80 is not a current-model estimate. In our study it motivates the fixed visual-text arm plus matched text-only and benign-OCR controls. [4] PolyJailbreak: use method/algorithm and Figure 6 for the iterative workflow; Table VI reports average ASR 83.34% and harmfulness 3.976/5 over eight models. The paper states maximum 15 optimization steps; do not conflate optimization steps with the common target-call cap proposed in our protocol. Log auxiliary judge/agent/image-generation calls and keep the attack's own metric separate. Neither paper alone validates our proposed training pipeline."
        },
        {
            "title": "Evaluation validity: visual leakage and broad risk coverage",
            "refs": "[2] VLSBench, §§2–4 / Fig. 1 / Table 4 | [3] USB, taxonomy / Figs. 2–3 / Tables 1–3",
            "left_title": "VLSBench | test whether the image matters [2]",
            "left_body": "CONTRIBUTION\nDiagnoses visual safety information leakage (VSIL): text can reveal the unsafe intent, letting a model pass without relying on the image. Dataset: 2,241 image-text pairs, 1,957 images, six categories, 19 subcategories.\n\nEVALUATION + RESULT\nTable 4 reports LLaVA-1.5-7B total safety (refusal + warning): multimodal SFT 21.26%, multimodal DPO 27.01%, text SFT 13.99%, text DPO 13.99%. This is not ASR.\n\nPROJECT MAPPING\nUse leakless cases and original/removed/masked/decoy image controls; score image dependence separately.",
            "left_role": "Pipeline: leak-resistant evaluation + causal image controls",
            "right_title": "USB | unified safety evaluation [3]",
            "right_body": "CONTRIBUTION\nCrosses 61 risk categories with four image/text risk states (244 intersections); USB-Base has 13,175 samples and USB-Hard 3,785. Evaluates 22 MLLMs with safety and refusal outcomes.\n\nEVALUATION + RESULT\nTable 2 reports Claude Sonnet 4 at SR 91.16% and RR 18.30% (table-reported RR variation ±0.75). Preserve the paper's metric definitions and directions.\n\nPROJECT MAPPING\nSelect relevant VLM image/text strata; report unsafe compliance and benign false refusal by cell, not one pooled score. USB is a benchmark, not an attack or training recipe.",
            "right_role": "Pipeline: benchmark strata + safety/over-refusal reporting",
            "notes": "[2] VLSBench is an evaluation-validity contribution, not a defense. Explain VSIL: unsafe intent can be explicit in text so the image is not needed. The dataset is 2,241 pairs/1,957 images, six categories/19 subcategories. Table 4 calls its outcome total safety rate and combines refusal and warning; for LLaVA-1.5-7B the cited values are 21.26% (multimodal SFT), 27.01% (multimodal DPO), 13.99% (text SFT), and 13.99% (text DPO). Do not rename this ASR. Our image ablations should use matched cases and be interpreted as attribution controls. [3] USB's taxonomy crosses 61 categories with four risk states and supplies Base/Hard resources; paper reports 22 MLLMs. The cited Table 2 Claude Sonnet 4 pair is SR 91.16%, RR 18.30% (with table variation for RR). Keep the paper's precise metric definitions, split, and table caption in view. Project use is selected relevant strata, not wholesale pooling, and paired benign controls."
        },
        {
            "title": "Training resources: supervised safety data and preferences",
            "refs": "[5] VLGuard, §3 / §4 / Table 2 / Fig. 3 | [6] SPA-VL, dataset construction / §4 / Table 2 / Fig. 1",
            "left_title": "VLGuard | multimodal safety SFT baseline [5]",
            "left_body": "CONTRIBUTION\nA low-cost multimodal safety fine-tuning resource: 2,000 train images (977 harmful, 1,023 safe), about 3,000 pairs; 1,000 test images across Safe-Safe, Safe-Unsafe, and Unsafe conditions.\n\nEVALUATION + RESULT\nCompares post-hoc and mixed fine-tuning. Table 2 shows reduced attack/unsafe-condition rates; preserving ordinary helpfulness examples matters for benign performance. Read each column's direction before quoting it.\n\nPROJECT MAPPING\nUse as the multimodal SFT baseline with a matched helpfulness mixture and safe-neighbor evaluation; split by source intent first.",
            "left_role": "Pipeline: image-conditioned safety SFT baseline",
            "right_title": "SPA-VL | preference data + DPO/PPO [6]",
            "right_body": "CONTRIBUTION\n100,788 preference records from 12 response models; each is (question, image, chosen answer, rejected answer). DPO/PPO are optimization methods, not datasets.\n\nEVALUATION + RESULT\nTable 2 tests DPO/PPO on MM-SafetyBench, AdvBench, HarmEval, and HelpEval. SPA-VL-DPO reports 0.60% average MM-SafetyBench ASR and 0.00/0.00 on AdvBench vanilla/suffix under that paper's setup.\n\nPROJECT MAPPING\nUse the resource provenance and standard DPO comparator; audit source/intent overlap and retain a separately held-out attack family.",
            "right_role": "Pipeline: preference data + standard DPO comparator",
            "notes": "[5] VLGuard's contribution is a small multimodal safety training resource and SFT comparison, not a universal defense. Its 2,000 training images (977 harmful/1,023 safe), about 3,000 image-instruction-response pairs, and 1,000-image test set are reported in the paper. The test conditions are Safe-Safe, Safe-Unsafe, and Unsafe. Table 2 compares baseline, post-hoc, and mixed-finetuning conditions across text and image safety evaluations; use the exact column headings and arrows. The result is lower unsafe response/attack rates, with helpfulness examples relevant to over-refusal. Our role: minimum multimodal SFT arm, matched helpfulness mixture, safe-neighbor false-refusal checks. [6] SPA-VL reports 100,788 preference examples, 12 source response models, and a split of 93,258 train, 7,000 validation, 265 HarmEval, 265 HelpEval. A record contains question, image, chosen, rejected response. Table 2 reports DPO/PPO outcomes on several benchmarks; the listed DPO values are 0.60 average MM-SafetyBench ASR and 0.00/0.00 AdvBench vanilla/suffix under the authors' test protocol. DPO is a preference optimization algorithm, not a dataset. Check the paper's taxonomy-count inconsistency and split provenance before reuse. Our role: preference source and ordinary multimodal DPO comparator, with intent-level decontamination."
        },
        {
            "title": "Robust alignment and integrated visual safety",
            "refs": "[7] ADPO, §§3–4 / Tables 1–2 | [8] HoliSafe, §§2–4 / Tables 1, 3, 7 / Fig. 2",
            "left_title": "ADPO | adversarial preference alignment [7]",
            "left_body": "CONTRIBUTION\nAdversary-aware DPO adds image/latent adversarial examples to preference alignment; it is direct prior art for robustness-oriented VLM training.\n\nEVALUATION + RESULT\nFor LLaVA-1.5-7B, Table 1's five ASRs (VisualAdv / MMPGDBlank / MultiTrust typo / multimodal / crossmodal) are 5.0 / 0.5 / 0.0 / 0.0 / 0.2% with ADPO, versus 64.5 / 84.0 / 22.2 / 55.1 / 42.0% base. Utility shifts and training cost are also reported.\n\nPROJECT MAPPING\nCandidate comparator; reserve attack families/templates for test and count optimization compute.",
            "left_role": "Pipeline: adversarial preference-training comparator",
            "right_title": "HoliSafe / Safe-VLM | five-state coverage + VGM [8]",
            "right_body": "CONTRIBUTION\n14,246 image-instruction-response pairs; 4,031 HoliSafe-Bench questions; five image/text safety states. Safe-VLM adds a learnable safety token/visual guard and joint generation/classification objective.\n\nEVALUATION + RESULT\nTable 3 reports judge-dependent rates; Safe-LLaVA-7B has 8.8% mASR / 1.3% benign RR under Claude-3.5, and 15.3% mASR under GPT-4o in the cited arXiv version.\n\nPROJECT MAPPING\nReuse/adapt its coverage taxonomy; treat VGM as one separate comparator, not a stackable component. Report judge and version.",
            "right_role": "Pipeline: modality-state coverage + optional VGM arm",
            "notes": "[7] ADPO is the closest direct prior for adversarial multimodal preference alignment, so novelty must be framed as controlled transfer/generalization, not the first such method. Table 1 for LLaVA-1.5-7B gives base attack ASRs 64.5/84.0/22.2/55.1/42.0 and ADPO 5.0/0.5/0.0/0.0/0.2 in the stated attack-column order. The paper also evaluates utility and iteration cost; do not imply utility is unchanged. Project mapping: comparator only if training recipe/code and white-box assumptions are reproducible; hold out families/templates and account for perturbation compute. [8] HoliSafe joins benchmark/data and method. It has 14,246 pairs, 10,215 train, 4,031 HoliSafe-Bench examples on 1,796 images; five states are unsafe image+unsafe text, unsafe image+safe text, safe image+unsafe text, safe image+safe text yielding unsafe output, and safe output. Safe-VLM uses a safety meta token/visual guard with classification and generation objectives. Table 3 judge-specific Safe-LLaVA-7B result in the cited arXiv version: 8.8% mASR/1.3% RR under Claude-3.5-Sonnet and 15.3% mASR under GPT-4o; do not mix versions or substitute another model's 12.3%. Table 7 is an ablation, not the main result. Project mapping: cite as prior art for five-state framing and optional VGM arm; do not claim the taxonomy or guard as novel."
        },
        {
            "title": "Visual-side and safety–helpfulness alignment",
            "refs": "[9] DAVSP, §§3–5 / Figs. 1–2 / Tables 1–3 | [10] Pragma-VL, §§3–4 / Figs. 1–4 / Tables 2–3",
            "left_title": "DAVSP | visual prompt / activation alignment [9]",
            "left_body": "CONTRIBUTION\nOptimizes a visual safety prompt and activation alignment while keeping the base LVLM frozen; inference adds the learned visual and textual safety prompts.\n\nEVALUATION + RESULT\nTable 2 reports resistance success rate (RSR), not ASR: FigStep 84.20% on LLaVA-1.5-13B and 99.20% on Qwen2-VL-7B; MM-SafetyBench SD+TYPO 98.72% and 99.12%.\n\nPROJECT MAPPING\nSeparate visual-side comparator; include capability and latency controls. RSR is paper-specific and not directly comparable with ASR.",
            "left_role": "Pipeline: optional inference-time visual intervention",
            "right_title": "Pragma-VL | pragmatic arbitration [10]",
            "right_body": "CONTRIBUTION\nA staged training framework for contextual safety/helpfulness arbitration, using data augmentation and reward modeling; not a jailbreak attack or simple SFT dataset.\n\nEVALUATION + RESULT\nQwen2.5-VL-7B Table 2: MM-SafetyBench ASR 31.66% vs. base 48.75%; SIUO safety 63.47% vs. 38.78%, with effectiveness 95.21% vs. 92.17%. Table 3 also checks VQA/math capability.\n\nPROJECT MAPPING\nPotential training comparator if recipe, data, and compute are feasible; report safety, helpfulness, capability separately.",
            "right_role": "Pipeline: optional training-time comparator / trade-off evidence",
            "notes": "[9] DAVSP is distinct from weight-update SFT/DPO: it optimizes a visual safety prompt and deep alignment in activation space while freezing the base model, then deploys a padding/border prompt together with a text safety prompt. Table 2 calls the metric RSR; higher is more resistant. Cited values are 84.20%/99.20% on FigStep for LLaVA-1.5-13B/Qwen2-VL-7B and 98.72%/99.12% on MM-SafetyBench SD+TYPO. Table 1 capability results and other utility results should accompany any preservation claim. Project role: optional independent inference intervention, with latency/utility measured; never label RSR as ASR. [10] Pragma-VL is a staged context-aware safety/helpfulness arbitration method. In Table 2 for Qwen2.5-VL-7B, MM-SafetyBench ASR is 31.66% vs 48.75% base; SIUO effectiveness/safety 95.21/63.47 vs 92.17/38.78; MSSBench effectiveness/safety 99.66/55.89 vs 98.48/36.53. Table 3 has VQAv2 and MathVista capability results, including 84.20% vs 83.60% VQAv2 and 67.20% vs 67.80% MathVista. Interpret each metric using its paper definition; this is not a cross-paper ranking. Project role: recent feasible-method comparator only after code/data/compute feasibility; otherwise use as context for separate safety/helpfulness/capability outcomes."
        },
    ]
    added = [add_review_slide(prs, item, 5 + i) for i, item in enumerate(review_slides)]
    attack_slide = add_attack_coverage_slide(prs, 10)

    # Place the new evidence section after the existing literature map, preserving all existing slides.
    order = prs.slides._sldIdLst
    new_elements = list(order)[-6:]
    for element in new_elements:
        order.remove(element)
    for offset, element in enumerate(new_elements):
        order.insert(6 + offset, element)

    # Shift printed page labels on the original content slides after the inserted section.
    for slide in list(prs.slides)[12:]:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.top > Inches(6.8) and shape.left > Inches(11.5):
                label = shape.text.strip()
                if label.isdigit() and 5 <= int(label) <= 16:
                    shape.text = f"{int(label) + 6:02d}"

    # Keep the benchmark/data page explicit about resource roles and the attack page honest about scope.
    old_benchmark = prs.slides[12]
    for shape in old_benchmark.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip() == "What the core test resources actually cover":
            shape.text = "Separate training data from locked evaluation benchmarks"
    old_benchmark.notes_slide.notes_text_frame.text += (
        "\n\nResource-role reminder: VLGuard and SPA-VL are principally training resources (safety SFT and preference data). "
        "VLSBench, USB, HoliSafe-Bench, SIUO, SafeBench/FigStep, and selected JailBreakV-28K items are evaluation resources with "
        "different scopes and metric definitions. No single one replaces the others: use leakless/image-dependent cases, broad risk/modality "
        "strata, compositional safe-input/unsafe-output cases, and benign near-neighbors. Split source intents before deriving any variants."
    )

    prs.save(PPTX)
    check = Presentation(PPTX)
    if len(check.slides) != 24:
        raise RuntimeError("Unexpected slide count after adding the literature section.")
    if fingerprint(check.slides[1]) != before_slide2 or fingerprint(check.slides[2]) != before_slide3:
        raise RuntimeError("Slide 2 or blank slide 3 changed unexpectedly.")
    expected = ["Attacks: from fixed visual prompts to adaptive search",
                "Evaluation validity: visual leakage and broad risk coverage",
                "Training resources: supervised safety data and preferences",
                "Robust alignment and integrated visual safety",
                "Visual-side and safety–helpfulness alignment",
                "A representative suite, not every published VLM attack"]
    for i, title in enumerate(expected, 6):
        if not any(getattr(s, "has_text_frame", False) and s.text.strip() == title for s in check.slides[i].shapes):
            raise RuntimeError(f"Missing inserted slide: {title}")
    print(f"Added five paper-by-paper evidence slides and one attack coverage slide to {PPTX}.")
    print(f"Verified slide 2 ({sum(s.shape_type == 13 for s in check.slides[1].shapes)} saved images) is unchanged, slide 3 remains blank, and the deck now has 24 slides.")


if __name__ == "__main__":
    main()
