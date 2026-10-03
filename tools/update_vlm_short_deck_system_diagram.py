from hashlib import sha256
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


PPTX = Path(__file__).resolve().parents[1] / "VLM_Safety_Alignment_10_Paper_Proposal.pptx"
INK = RGBColor(30, 30, 30)
GRAY = RGBColor(242, 242, 242)
WHITE = RGBColor(255, 255, 255)


def slide_fingerprint(slide):
    values = []
    for shape in slide.shapes:
        item = [shape.shape_type, shape.left, shape.top, shape.width, shape.height]
        if getattr(shape, "has_text_frame", False):
            item.append(shape.text)
        if shape.shape_type == 13:
            item.append(sha256(shape.image.blob).hexdigest())
        values.append(tuple(item))
    return tuple(values)


def add_text(slide, x, y, w, h, value, size, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(.04)
    frame.margin_right = Inches(.04)
    frame.margin_top = Inches(.025)
    frame.margin_bottom = Inches(.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE if align == PP_ALIGN.CENTER else MSO_ANCHOR.TOP
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = value
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = INK
    return box


def add_stage(slide, x, width, heading, body):
    y, height, header_h = 1.72, 3.58, .55
    body_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    body_shape.fill.solid()
    body_shape.fill.fore_color.rgb = WHITE
    body_shape.line.color.rgb = INK
    body_shape.line.width = Pt(1.25)

    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(header_h)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = GRAY
    header.line.color.rgb = INK
    header.line.width = Pt(1.0)

    add_text(slide, x + .06, y + .06, width - .12, .4, heading, 12, True, PP_ALIGN.CENTER)
    add_text(slide, x + .12, y + .68, width - .24, height - .8, body, 10.5)
    return x, y, width, height


def main():
    prs = Presentation(PPTX)
    if len(prs.slides) != 18:
        raise RuntimeError(f"Expected saved 18-slide deck; found {len(prs.slides)}. No changes made.")

    slide2_before = slide_fingerprint(prs.slides[1])
    slide3_before = slide_fingerprint(prs.slides[2])
    if sum(shape.shape_type == 13 for shape in prs.slides[1].shapes) != 4:
        raise RuntimeError("Slide 2 no longer has the verified four images. No changes made.")
    if len(prs.slides[2].shapes) != 0:
        raise RuntimeError("Slide 3 is no longer empty. No changes made.")

    matches = [slide for slide in prs.slides if any(
        getattr(shape, "has_text_frame", False)
        and shape.text.strip() == "Proposed end-to-end study design"
        for shape in slide.shapes
    )]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one system-design slide; found {len(matches)}. No changes made.")

    slide = matches[0]
    if not any(getattr(shape, "has_text_frame", False) and "Source intents" in shape.text
               for shape in slide.shapes):
        raise RuntimeError("System-design slide differs from the expected prior layout. No changes made.")

    # Preserve the slide title, rule, citation footer, and page label; replace only diagram content.
    for shape in list(slide.shapes)[5:]:
        shape._element.getparent().remove(shape._element)

    stages = [
        add_stage(slide, .64, 2.02, "1 | DATA + EVIDENCE",
                  "VLGuard: safety examples [5]\nSPA-VL: preference pairs [6]\nHoliSafe: image/text cases [8]\nVLSBench + USB: evaluation design [2,3]\n\nTrack source, license, intent, and safe counterpart."),
        add_stage(slide, 2.86, 2.10, "2 | LABEL + SPLIT",
                  "Label image, text, and joint risk; add safe near-neighbors.\n\nSplit by source intent BEFORE paraphrase, rendering, image synthesis, or attack transforms.\n\nTrain / validation / locked test; deduplicate intents and images."),
        add_stage(slide, 5.16, 2.58, "3 | SEPARATE ALIGNMENT ARMS",
                  "B0 Native checkpoint\nB1 Text-only safety SFT\nB2 Multimodal safety SFT (VLGuard)\nB3 Standard DPO (SPA-VL)\nB4 ADPO (if reproducible)\n\nChoose at most one feasible comparator: HoliSafe VGM, DAVSP, or Pragma-VL. Do not stack methods."),
        add_stage(slide, 7.94, 2.62, "4 | LOCKED EVALUATION",
                  "ATTACKS\nFixed: FigStep [1], HADES +Opt [11]\nAttack set: JailBreakV [12]\nAdaptive: PolyJailbreak [4]\n\nBENCHMARKS + CONTROLS\nVLSBench / USB / HoliSafe [2,3,8]\nSafe neighbors; image removed, masked, or decoy. Same frozen inference settings."),
        add_stage(slide, 10.76, 1.94, "5 | SCORE + REPORT",
                  "Harmful compliance / ASR\nBenign false refusal\nSafe-answer quality\nImage grounding\nVQA / OCR capability\nCalls, latency, compute\n\nPaired by intent; CIs + human audit."),
    ]

    for left, right in zip(stages, stages[1:]):
        x1 = left[0] + left[2] + .015
        x2 = right[0] - .015
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(3.48), Inches(x2), Inches(3.48)
        )
        connector.line.color.rgb = INK
        connector.line.width = Pt(1.4)
        connector.line.end_arrowhead = True

    add_text(slide, .78, 5.48, 11.75, .38,
             "Locked test intents never enter training. Attacks are procedures; benchmarks are test sets; image ablations test dependence.",
             12, True, PP_ALIGN.CENTER)
    add_text(slide, .78, 5.94, 11.75, .42,
             "Compare interventions separately. Preserve each metric, judge/version, query budget, and cost; this is not a cross-paper leaderboard.",
             11, False, PP_ALIGN.CENTER)

    slide.notes_slide.notes_text_frame.text = (
        "This is the proposed system design, synthesized from the literature rather than a pipeline already validated by one paper. "
        "Stage 1 distinguishes training resources from evaluation resources. Stage 2 splits source intents before generating any derived image, "
        "paraphrase, or attack variant, preventing variants of one intent from crossing train and test. Stage 3 compares separate arms: "
        "native model, text-only safety SFT, multimodal safety SFT, ordinary preference tuning, ADPO if faithfully reproducible, and at most "
        "one recent comparator after feasibility checks. Stage 4 applies locked fixed and adaptive attacks across benchmark strata and includes "
        "benign controls and image ablations. Stage 5 reports harmful compliance, benign false refusal, safe-answer quality, visual grounding, "
        "general capability, and cost separately, paired by source intent with confidence intervals and human audit. The shared query budget is "
        "a proposed protocol choice; PolyJailbreak reports 15 optimization steps, not necessarily 15 total API calls. Slides 2 and 3 were preserved."
    )

    prs.save(PPTX)
    check = Presentation(PPTX)
    if len(check.slides) != 18:
        raise RuntimeError("Slide count changed unexpectedly after save.")
    if slide_fingerprint(check.slides[1]) != slide2_before:
        raise RuntimeError("Slide 2 changed unexpectedly during update.")
    if slide_fingerprint(check.slides[2]) != slide3_before or len(check.slides[2].shapes) != 0:
        raise RuntimeError("Slide 3 changed unexpectedly during update.")
    if not any(getattr(shape, "has_text_frame", False) and "B0 Native checkpoint" in shape.text
               for shape in check.slides[12].shapes):
        raise RuntimeError("Updated diagram was not found on slide 13.")
    print(f"Updated diagram in {PPTX}; slide 2 image fingerprint and blank slide 3 verified unchanged.")


if __name__ == "__main__":
    main()
